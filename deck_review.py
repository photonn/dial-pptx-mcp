"""
Deck-level review: what the deck *says*, as opposed to how each slide looks.

The vision reviewer judges slides a batch at a time, so anything that only
shows up across slides is invisible to it: an agenda on slide 2 promising a
section the deck never delivers, a section divider whose number disagrees with
the agenda, one slide's cards landing on the next slide's divider while their
own slide stays empty, the same KPI quoted with two different values. Those are
the defects an orchestrator building slide by slide actually produces — an
off-by-one `slide_index` moves content, it does not break a layout — and no
amount of per-slide geometry review catches them.

So the whole deck is also reviewed as text: `slide_outline` extracts every
slide's layout, title, text, charts, tables and pictures (plus which text is
drawn over by a later shape — the "content hidden behind a card" case), and one
text-only model call checks it against `COHERENCE_PROMPT` (run by
visual_qa._review, in parallel with the vision batches). It is cheap: no
rendering, no images, one call however long the deck is.

The outline is also what the vision reviewer receives beside each image (see
visual_qa.review_prompt), so it can tell text that exists but is not visible
from text that was never there.
"""
import json
import os

from pptx.enum.dml import MSO_FILL
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER
from pptx.oxml.ns import qn
from pptx.util import Emu


# Per-slide and per-element text budgets: enough for the reviewer to recognise
# what a slide is about, bounded so a 60-slide deck stays a modest prompt.
ELEMENT_TEXT_CHARS = 300
SLIDE_TEXT_CHARS = 1500
# Share of a text shape's area a later (higher z-order) shape must cover
# before the text is reported as drawn over.
COVERED_FRACTION = 0.3

# Placeholders that are slide furniture, not the slide's content.
_FURNITURE = {PP_PLACEHOLDER.FOOTER, PP_PLACEHOLDER.SLIDE_NUMBER,
              PP_PLACEHOLDER.DATE, PP_PLACEHOLDER.HEADER}
_TITLES = {PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE,
           PP_PLACEHOLDER.VERTICAL_TITLE}


def coherence_enabled() -> bool:
    """VISUAL_QA_COHERENCE (default true) switches the deck-level text review
    off, e.g. when the vision model is too slow for one more call."""
    return os.environ.get("VISUAL_QA_COHERENCE", "true").strip().lower() \
        not in ("0", "false", "no", "off")


def _in(v):
    return round(Emu(v).inches, 2) if v is not None else None


def _box(shape):
    try:
        return (shape.left, shape.top, shape.width, shape.height)
    except Exception:  # a shape with no resolvable geometry
        return None


def _overlap_fraction(inner, outer):
    """Share of `inner`'s area that `outer` covers (both EMU boxes)."""
    if not inner or not outer or None in inner or None in outer:
        return 0.0
    left = max(inner[0], outer[0])
    top = max(inner[1], outer[1])
    right = min(inner[0] + inner[2], outer[0] + outer[2])
    bottom = min(inner[1] + inner[3], outer[1] + outer[3])
    area = inner[2] * inner[3]
    if right <= left or bottom <= top or area <= 0:
        return 0.0
    return (right - left) * (bottom - top) / area


def _is_filled(shape):
    """Whether a shape paints an opaque area that can hide what is under it:
    pictures, charts, tables, groups, and autoshapes/text boxes with a fill."""
    if getattr(shape, "has_chart", False) or getattr(shape, "has_table", False):
        return True
    shape_type = shape.shape_type
    if shape_type in (MSO_SHAPE_TYPE.PICTURE, MSO_SHAPE_TYPE.GROUP):
        return True
    try:
        fill_type = shape.fill.type
    except Exception:  # connectors, graphic frames: no fill API
        return False
    if fill_type is not None:
        # python-pptx reports <a:noFill/> as BACKGROUND.
        return fill_type != MSO_FILL.BACKGROUND
    # No explicit fill: an autoshape still paints its theme style's fill
    # (add_shape writes <p:style><a:fillRef idx="…">); a text box has none.
    fill_ref = shape._element.find(".//" + qn("a:fillRef"))
    return fill_ref is not None and fill_ref.get("idx", "0") != "0"


def drawn_over(shapes):
    """{index: [later shape indexes whose opaque area covers it]} for every
    text-bearing shape. Later shapes are drawn on top, so a filled card added
    after a slide's title hides the title even though both "exist"."""
    boxes = [_box(s) for s in shapes]
    result = {}
    for i, shape in enumerate(shapes):
        if not (shape.has_text_frame and shape.text_frame.text.strip()):
            continue
        covering = [j for j in range(i + 1, len(shapes))
                    if _is_filled(shapes[j])
                    and _overlap_fraction(boxes[i], boxes[j]) >= COVERED_FRACTION]
        if covering:
            result[i] = covering
    return result


def _kind(shape):
    if shape.is_placeholder:
        try:
            return "placeholder:" + shape.placeholder_format.type.name.lower()
        except Exception:
            return "placeholder"
    if getattr(shape, "has_chart", False):
        return "chart"
    if getattr(shape, "has_table", False):
        return "table"
    shape_type = shape.shape_type
    if shape_type == MSO_SHAPE_TYPE.PICTURE:
        return "picture"
    if shape_type == MSO_SHAPE_TYPE.GROUP:
        return "group"
    if shape_type == MSO_SHAPE_TYPE.TEXT_BOX:
        return "text_box"
    return "shape"


def _clip(text, limit):
    text = " / ".join(line.strip() for line in text.splitlines() if line.strip())
    return text if len(text) <= limit else text[:limit] + "…"


def _group_text(group):
    parts = []
    for member in group.shapes:
        if member.has_text_frame and member.text_frame.text.strip():
            parts.append(member.text_frame.text)
        elif member.shape_type == MSO_SHAPE_TYPE.GROUP:
            parts.append(_group_text(member))
    return " / ".join(p for p in parts if p)


def _chart_summary(chart):
    info = {}
    try:
        info["chart_type"] = chart.chart_type.name.lower()
    except Exception:
        pass
    try:
        if chart.has_title:
            info["title"] = chart.chart_title.text_frame.text[:80]
    except Exception:
        pass
    try:
        info["series"] = [s.name for s in chart.series][:8]
        info["categories"] = [str(c)[:30] for c in chart.plots[0].categories][:12]
    except Exception:
        pass
    return info


def _is_footnote(shape, slide_height):
    """A short text line in the bottom band — a source line, not content."""
    box = _box(shape)
    if not box or None in box or not slide_height:
        return False
    return box[1] > slide_height * 0.85 and box[3] < slide_height * 0.1


def slide_outline(pres, slide_numbers=None):
    """Structured text outline of the deck's slides (1-based numbers; None
    means every slide) — see the module docstring for what it is for."""
    outline = []
    slide_height = pres.slide_height
    for number, slide in enumerate(pres.slides, start=1):
        if slide_numbers and number not in slide_numbers:
            continue
        shapes = list(slide.shapes)
        covered = drawn_over(shapes)
        title, elements, body, budget = None, [], 0, SLIDE_TEXT_CHARS
        for index, shape in enumerate(shapes):
            kind = _kind(shape)
            entry = {"shape_index": index, "kind": kind}
            is_title = is_furniture = False
            if shape.is_placeholder:
                try:
                    ph_type = shape.placeholder_format.type
                    is_title = ph_type in _TITLES
                    is_furniture = ph_type in _FURNITURE
                except Exception:
                    pass
            text = ""
            if shape.has_text_frame:
                text = shape.text_frame.text
            elif kind == "group":
                text = _group_text(shape)
            if kind == "table":
                table = shape.table
                entry["table"] = {
                    "rows": len(table.rows), "columns": len(table.columns),
                    "header": [table.cell(0, c).text_frame.text[:40]
                               for c in range(len(table.columns))]}
                text = " | ".join(
                    table.cell(r, c).text_frame.text
                    for r in range(len(table.rows))
                    for c in range(len(table.columns)))
            if kind == "chart":
                entry["chart"] = _chart_summary(shape.chart)
            if text.strip():
                clipped = _clip(text, min(ELEMENT_TEXT_CHARS, max(budget, 40)))
                budget -= len(clipped)
                entry["text"] = clipped
            elif shape.has_text_frame and kind.startswith("placeholder"):
                entry["empty"] = True
            if index in covered:
                entry["drawn_over_by"] = covered[index]
            box = _box(shape)
            if box and None not in box:
                entry["box_in"] = [_in(v) for v in box]
            if is_title and text.strip() and title is None:
                title = _clip(text, 200)
                entry["role"] = "title"
            elif is_furniture:
                entry["role"] = "furniture"
            elif text.strip() and _is_footnote(shape, slide_height):
                entry["role"] = "footnote"
            elif text.strip() or kind in ("picture", "chart", "table", "group"):
                body += 1
            elements.append(entry)
        outline.append({
            "slide": number,
            "layout": slide.slide_layout.name,
            "title": title,
            "body_elements": body,
            "elements": elements,
        })
    return outline


COHERENCE_PROMPT = """You are reviewing the STORY and STRUCTURE of a \
PowerPoint deck, not its looks. Below is a text outline of every slide: its \
layout name, title, and each element (shape_index, kind, text, box in inches \
[left, top, width, height], and "drawn_over_by" when later shapes cover it). \
"body_elements" counts the slide's content besides its title, footers and \
source lines.

{outline}

Find every problem of these kinds:

1. Agenda / index consistency. If a slide lists the deck's sections (an \
agenda, contents, index or overview — usually slide 2), compare it with the \
section divider slides and the slide titles that follow: every agenda item \
must have a matching section in the same order and with the same numbering; \
every section divider must appear in the agenda; wording should match closely \
enough that a reader recognises the section. Report missing sections, extra \
sections, wrong order, wrong numbers and wording mismatches. Name the agenda \
slide in "slide" and the section slides in "related_slides".
2. Content that does not match its slide. A slide's body must be about what \
its title says. Content that clearly belongs to ANOTHER slide's title (the \
typical symptom of content placed with an off-by-one slide number: one slide \
empty, its neighbour carrying its content, or a section divider carrying \
cards) is critical — report it on the slide that holds the misplaced content, \
with the slide it belongs on in "related_slides", and name the shape_index of \
each misplaced element in "element".
3. Empty or unfinished slides. A content slide (not a title, section divider, \
statement, quote or closing slide) with body_elements 0 is critical: it will \
be presented blank. Also report empty placeholders that leave a visible hole, \
placeholder or sample text ("Click to add", "Lorem ipsum", "TBD", "XXX", \
"[insert …]"), and a title with no content that belongs to it.
4. Hidden content. Text that is "drawn_over_by" an opaque shape is very likely \
invisible (e.g. a divider's title covered by cards). Report it.
5. Duplicates and contradictions. Duplicate slides or near-identical \
content; the same fact, figure, date or name given differently on two slides \
(e.g. €961m on one slide, €960m on another); inconsistent terminology or \
spelling of the same product, unit or name; a percentage or total that does \
not add up where the numbers are all on the slides.
6. Narrative order. Section dividers followed by slides of another section; \
the closing / thank-you slide not near the end; appendix material before the \
main content; a summary that cites points the deck never makes.
7. Titles. Inconsistent title style across comparable slides (sentence vs \
title case, trailing punctuation), a title that is truncated or ends \
mid-sentence, a numbered title whose number is out of sequence.
8. Text quality. Obvious typos, doubled words, sentences cut off mid-word, \
typed bullet characters or numbering duplicated inside text ("• •", "1. 1."), \
literal markup or escape sequences ("\\\\n", "**", "&amp;"), mojibake \
("â€™") or missing-glyph boxes.

Severity: "critical" when a slide would be presented wrong (blank, \
misplaced or hidden content, agenda promising what the deck does not \
deliver); "major" for anything an audience would notice (contradictory \
figures, wrong order or numbering, visible placeholder text, typos in \
titles); "minor" for polish.

Respond with ONLY a JSON object, no markdown fence:
{{"passed": true|false, "issues": [{{"slide": <1-based number>, "severity": \
"critical"|"major"|"minor", "category": "agenda"|"misplaced_content"|\
"empty_slide"|"hidden_content"|"duplicate"|"contradiction"|"order"|"title"|\
"text_quality", "element": "<shape_index(es) or element name>", \
"related_slides": [<1-based numbers>], "description": "...", \
"suggested_fix": "..."}}]}}
In "suggested_fix" be concrete: the exact corrected agenda text, which \
shape_index moves to which slide, what the empty slide needs. Never invent \
facts or figures — when content is missing, say what the slide needs, not \
what it should claim. "passed" is true only when there are no critical or \
major issues."""


def coherence_prompt(outline):
    return COHERENCE_PROMPT.format(
        outline=json.dumps(outline, ensure_ascii=False, separators=(",", ":")))

