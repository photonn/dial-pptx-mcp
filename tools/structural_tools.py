"""
Structural element tools for PowerPoint MCP Server.
Handles tables, shapes, and charts.
"""
from typing import Dict, List, Optional, Any
from mcp.server.fastmcp import FastMCP
from mcp.types import ToolAnnotations
from pptx.enum.text import MSO_ANCHOR
from pptx.oxml.ns import qn
from pptx.util import Inches
import utils as ppt_utils
from logging_utils import get_logger

logger = get_logger("tools.structural")

def _drop_shadow(shape):
    """No shadow: an empty a:effectLst on the shape, and the style's
    effectRef pointed at the theme's empty effect style. The first alone is
    what the spec says is enough; LibreOffice still draws the style's shadow,
    and visual QA renders through LibreOffice."""
    shape.shadow.inherit = False
    style = shape._element.find(qn("p:style"))
    if style is not None:
        ref = style.find(qn("a:effectRef"))
        if ref is not None:
            ref.set("idx", "0")


_ANCHORS = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE,
            "bottom": MSO_ANCHOR.BOTTOM}


def register_structural_tools(app: FastMCP, presentations: Dict, get_current_presentation_id, validate_parameters, is_positive, is_non_negative, is_in_range, is_valid_rgb, add_shape_direct):
    """Register structural element tools with the FastMCP app"""
    
    @app.tool(
        annotations=ToolAnnotations(
            title="Add Table",
        ),
    )
    def add_table(
        slide_index: int,
        rows: int,
        cols: int,
        left: float,
        top: float,
        width: float,
        height: float,
        data: Optional[List[List[str]]] = None,
        header_row: bool = True,
        header_font_size: int = 12,
        body_font_size: int = 10,
        header_bg_color: Optional[List[int]] = None,
        body_bg_color: Optional[List[int]] = None,
        border_color: Optional[List[int]] = None,
        presentation_id: Optional[str] = None
    ) -> Dict:
        """Add a table to a slide with enhanced formatting options."""
        pres_id = presentation_id if presentation_id is not None else get_current_presentation_id()
        
        if pres_id is None or pres_id not in presentations:
            return {
                "error": "Unknown or expired presentation_id. Pass the presentation_id returned by create_presentation, create_presentation_from_template, or open_presentation"
            }
        
        pres = presentations[pres_id]
        
        if slide_index < 0 or slide_index >= len(pres.slides):
            return {
                "error": f"Invalid slide index: {slide_index}. Available slides: 0-{len(pres.slides) - 1}"
            }
        
        slide = pres.slides[slide_index]
        
        # Validate parameters
        validations = {
            "rows": (rows, [(is_positive, "must be a positive integer")]),
            "cols": (cols, [(is_positive, "must be a positive integer")]),
            "left": (left, [(is_non_negative, "must be non-negative")]),
            "top": (top, [(is_non_negative, "must be non-negative")]),
            "width": (width, [(is_positive, "must be positive")]),
            "height": (height, [(is_positive, "must be positive")])
        }
        
        if header_bg_color is not None:
            validations["header_bg_color"] = (header_bg_color, [(is_valid_rgb, "must be a valid RGB list [R, G, B] with values 0-255")])
        if body_bg_color is not None:
            validations["body_bg_color"] = (body_bg_color, [(is_valid_rgb, "must be a valid RGB list [R, G, B] with values 0-255")])
        if border_color is not None:
            validations["border_color"] = (border_color, [(is_valid_rgb, "must be a valid RGB list [R, G, B] with values 0-255")])
        
        valid, error = validate_parameters(validations)
        if not valid:
            return {"error": error}
        
        # Validate data if provided
        if data:
            if len(data) != rows:
                return {
                    "error": f"Data has {len(data)} rows but table should have {rows} rows"
                }
            for i, row in enumerate(data):
                if len(row) != cols:
                    return {
                        "error": f"Row {i} has {len(row)} columns but table should have {cols} columns"
                    }
        
        try:
            # Add the table
            table_shape = ppt_utils.add_table(slide, rows, cols, left, top, width, height)
            table = table_shape.table
            
            # Populate with data if provided
            if data:
                for r in range(rows):
                    for c in range(cols):
                        if r < len(data) and c < len(data[r]):
                            table.cell(r, c).text = str(data[r][c])
            
            # Apply formatting
            for r in range(rows):
                for c in range(cols):
                    cell = table.cell(r, c)
                    
                    # Header row formatting
                    if r == 0 and header_row:
                        if header_bg_color:
                            ppt_utils.format_table_cell(
                                cell, bg_color=tuple(header_bg_color), font_size=header_font_size, bold=True
                            )
                        else:
                            ppt_utils.format_table_cell(cell, font_size=header_font_size, bold=True)
                    else:
                        # Body cell formatting
                        if body_bg_color:
                            ppt_utils.format_table_cell(
                                cell, bg_color=tuple(body_bg_color), font_size=body_font_size
                            )
                        else:
                            ppt_utils.format_table_cell(cell, font_size=body_font_size)
            
            logger.debug("table_added slide=%s rows=%s cols=%s data_rows=%s",
                         slide_index, rows, cols, len(data) if data else 0)
            return {
                "message": f"Added {rows}x{cols} table to slide {slide_index}",
                "shape_index": len(slide.shapes) - 1,
                "rows": rows,
                "cols": cols
            }
        except Exception as e:
            return {
                "error": f"Failed to add table: {str(e)}"
            }

    @app.tool(
        annotations=ToolAnnotations(
            title="Format Table Cell",
        ),
    )
    def format_table_cell(
        slide_index: int,
        shape_index: int,
        row: int,
        col: int,
        font_size: Optional[int] = None,
        font_name: Optional[str] = None,
        bold: Optional[bool] = None,
        italic: Optional[bool] = None,
        color: Optional[List[int]] = None,
        bg_color: Optional[List[int]] = None,
        alignment: Optional[str] = None,
        vertical_alignment: Optional[str] = None,
        presentation_id: Optional[str] = None
    ) -> Dict:
        """Format a specific table cell."""
        pres_id = presentation_id if presentation_id is not None else get_current_presentation_id()
        
        if pres_id is None or pres_id not in presentations:
            return {
                "error": "Unknown or expired presentation_id. Pass the presentation_id returned by create_presentation, create_presentation_from_template, or open_presentation"
            }
        
        pres = presentations[pres_id]
        
        if slide_index < 0 or slide_index >= len(pres.slides):
            return {
                "error": f"Invalid slide index: {slide_index}. Available slides: 0-{len(pres.slides) - 1}"
            }
        
        slide = pres.slides[slide_index]
        
        if shape_index < 0 or shape_index >= len(slide.shapes):
            return {
                "error": f"Invalid shape index: {shape_index}. Available shapes: 0-{len(slide.shapes) - 1}"
            }
        
        shape = slide.shapes[shape_index]
        
        try:
            if not hasattr(shape, 'table'):
                return {
                    "error": f"Shape at index {shape_index} is not a table"
                }
            
            table = shape.table
            
            if row < 0 or row >= len(table.rows):
                return {
                    "error": f"Invalid row index: {row}. Available rows: 0-{len(table.rows) - 1}"
                }
                
            if col < 0 or col >= len(table.columns):
                return {
                    "error": f"Invalid column index: {col}. Available columns: 0-{len(table.columns) - 1}"
                }
            
            cell = table.cell(row, col)
            
            ppt_utils.format_table_cell(
                cell,
                font_size=font_size,
                font_name=font_name,
                bold=bold,
                italic=italic,
                color=tuple(color) if color else None,
                bg_color=tuple(bg_color) if bg_color else None,
                alignment=alignment,
                vertical_alignment=vertical_alignment
            )
            
            return {
                "message": f"Formatted cell at row {row}, column {col} in table at shape index {shape_index} on slide {slide_index}"
            }
        except Exception as e:
            return {
                "error": f"Failed to format table cell: {str(e)}"
            }

    @app.tool(
        annotations=ToolAnnotations(
            title="Add Shape",
        ),
    )
    def add_shape(
        slide_index: int,
        shape_type: str,
        left: float,
        top: float,
        width: float,
        height: float,
        fill_color: Optional[List[int]] = None,
        line_color: Optional[List[int]] = None,
        line_width: Optional[float] = None,
        text: Optional[str] = None,  # Add text to shape
        font_size: Optional[int] = None,
        font_color: Optional[List[int]] = None,
        font_name: Optional[str] = None,
        bold: Optional[bool] = None,
        italic: Optional[bool] = None,
        alignment: Optional[str] = None,
        vertical_alignment: Optional[str] = None,
        text_margin: Optional[float] = None,
        shadow: bool = False,
        no_fill: bool = False,
        presentation_id: Optional[str] = None
    ) -> Dict:
        """Add an auto shape to a slide with enhanced options.

        A filled shape is drawn flat: no outline unless line_color is given,
        and no shadow unless shadow=True — templates' default shape style
        adds both, which is what makes generated cards look unlike the
        template's own. text may hold several lines separated by "\n" (one
        paragraph each); style it with font_size, font_color, font_name,
        bold, italic, alignment ("left"/"center"/"right") and
        vertical_alignment ("top"/"middle"/"bottom"); text_margin is the
        inner padding in inches. This makes a finished card, tile, badge or
        header box in one call — no separate text box stacked on top.
        no_fill=True draws only the outline (give line_color and line_width):
        rings, frames and arcs over a coloured background."""
        pres_id = presentation_id if presentation_id is not None else get_current_presentation_id()
        
        if pres_id is None or pres_id not in presentations:
            return {
                "error": "Unknown or expired presentation_id. Pass the presentation_id returned by create_presentation, create_presentation_from_template, or open_presentation"
            }
        
        pres = presentations[pres_id]
        
        if slide_index < 0 or slide_index >= len(pres.slides):
            return {
                "error": f"Invalid slide index: {slide_index}. Available slides: 0-{len(pres.slides) - 1}"
            }
        
        slide = pres.slides[slide_index]
        
        try:
            # Use the direct implementation that bypasses the enum issues
            shape = add_shape_direct(slide, shape_type, left, top, width, height)
            
            # Format the shape if formatting options are provided
            if any([fill_color, line_color, line_width]):
                ppt_utils.format_shape(
                    shape,
                    fill_color=tuple(fill_color) if fill_color else None,
                    line_color=tuple(line_color) if line_color else None,
                    line_width=line_width
                )
            if no_fill:
                shape.fill.background()
            elif fill_color and not line_color:
                shape.line.fill.background()
            if not shadow:
                _drop_shadow(shape)
            
            # Add text to shape if provided
            if text and hasattr(shape, 'text_frame'):
                frame = shape.text_frame
                frame.text = text
                frame.word_wrap = True
                if text_margin is not None:
                    margin = Inches(max(0.0, text_margin))
                    frame.margin_left = frame.margin_right = margin
                    frame.margin_top = frame.margin_bottom = margin
                anchor = _ANCHORS.get((vertical_alignment or "").lower())
                if anchor is not None:
                    frame.vertical_anchor = anchor
                if any(v is not None for v in (font_size, font_color, font_name,
                                               bold, italic, alignment)):
                    ppt_utils.format_text(
                        frame,
                        font_size=font_size,
                        font_name=font_name,
                        bold=bold,
                        italic=italic,
                        color=tuple(font_color) if font_color else None,
                        alignment=(alignment or "").lower() or None
                    )
            
            return {
                "message": f"Added {shape_type} shape to slide {slide_index}",
                "shape_index": len(slide.shapes) - 1
            }
        except ValueError as e:
            return {
                "error": str(e)
            }
        except Exception as e:
            return {
                "error": f"Failed to add shape '{shape_type}': {str(e)}"
            }

    @app.tool(
        annotations=ToolAnnotations(
            title="Add Chart",
        ),
    )
    def add_chart(
        slide_index: int,
        chart_type: str,
        left: float,
        top: float,
        width: float,
        height: float,
        categories: List[str],
        series_names: List[str],
        series_values: List[List[float]],
        has_legend: bool = True,
        legend_position: str = "right",
        has_data_labels: bool = False,
        title: Optional[str] = None,
        x_axis_title: Optional[str] = None,
        y_axis_title: Optional[str] = None,
        color_scheme: Optional[str] = None,
        font_size: Optional[int] = None,
        number_format: Optional[str] = None,
        value_axis_min: Optional[float] = None,
        show_value_axis: bool = True,
        show_gridlines: bool = True,
        presentation_id: Optional[str] = None
    ) -> Dict:
        """Add a chart to a slide with comprehensive formatting options.

        font_size sets every chart text (ticks, data labels, legend) in pt.
        number_format is the data-label format, e.g. '#,##0', '0.0"%"',
        '"€"#,##0"m"'. Bar, column and area charts of non-negative data start
        their value axis at 0 unless value_axis_min says otherwise — an axis
        starting at 3,900 turns a 6% rise into a doubling. With data labels
        on, show_value_axis=False and show_gridlines=False give the clean
        labelled-bars look.

        x_axis_title labels the CATEGORY axis and y_axis_title labels the
        VALUE axis, whatever the chart's orientation — the names are not
        screen directions. On column, line and area charts the category axis
        is the horizontal one, so x/y read as expected. On a BAR chart the
        bars run sideways: the category axis is VERTICAL and the value axis
        HORIZONTAL, so x_axis_title appears on the vertical axis. Naming a
        bar chart's axes by where they look on screen is the usual way these
        two end up swapped.

        A 'scatter' chart is the exception to both: it has no category axis,
        so pass its x values (as numbers) in categories and the matching y
        values in series_values. Use 'line_markers' when the horizontal axis
        is really a list of labels.
        """
        pres_id = presentation_id if presentation_id is not None else get_current_presentation_id()
        
        if pres_id is None or pres_id not in presentations:
            return {
                "error": "Unknown or expired presentation_id. Pass the presentation_id returned by create_presentation, create_presentation_from_template, or open_presentation"
            }
        
        pres = presentations[pres_id]
        
        if slide_index < 0 or slide_index >= len(pres.slides):
            return {
                "error": f"Invalid slide index: {slide_index}. Available slides: 0-{len(pres.slides) - 1}"
            }
        
        slide = pres.slides[slide_index]
        
        # Validate chart type
        valid_chart_types = [
            'column', 'stacked_column', 'bar', 'stacked_bar', 'line', 
            'line_markers', 'pie', 'doughnut', 'area', 'stacked_area', 
            'scatter', 'radar', 'radar_markers'
        ]
        if chart_type.lower() not in valid_chart_types:
            return {
                "error": f"Invalid chart type: '{chart_type}'. Valid types are: {', '.join(valid_chart_types)}"
            }
        
        # Validate series data
        if len(series_names) != len(series_values):
            return {
                "error": f"Number of series names ({len(series_names)}) must match number of series values ({len(series_values)})"
            }
        
        if not categories:
            return {
                "error": "Categories list cannot be empty"
            }
        
        # Validate that all series have the same number of values as categories
        for i, values in enumerate(series_values):
            if len(values) != len(categories):
                return {
                    "error": f"Series '{series_names[i]}' has {len(values)} values but there are {len(categories)} categories"
                }
        
        # A scatter chart has no category axis: `categories` carries its x
        # values, so they have to be numbers. Checked here to report it as an
        # instruction rather than as a "Failed to add chart" exception string.
        if chart_type.lower() in ppt_utils.SCATTER_CHART_TYPES:
            try:
                ppt_utils.parse_scatter_x_values(categories)
            except ValueError as e:
                return {"error": str(e)}
        
        try:
            # Add the chart
            chart = ppt_utils.add_chart(
                slide, chart_type, left, top, width, height,
                categories, series_names, series_values
            )
            
            if chart is None:
                return {"error": "Failed to create chart"}
            
            # Format the chart
            ppt_utils.format_chart(
                chart,
                has_legend=has_legend,
                legend_position=legend_position,
                has_data_labels=has_data_labels,
                title=title,
                x_axis_title=x_axis_title,
                y_axis_title=y_axis_title,
                color_scheme=color_scheme
            )
            ppt_utils.style_chart_values(
                chart, chart_type, series_values, font_size=font_size,
                number_format=number_format, value_axis_min=value_axis_min,
                show_value_axis=show_value_axis,
                show_gridlines=show_gridlines)
            
            return {
                "message": f"Added {chart_type} chart to slide {slide_index}",
                "shape_index": len(slide.shapes) - 1,
                "chart_type": chart_type,
                "series_count": len(series_names),
                "categories_count": len(categories)
            }
        except Exception as e:
            return {
                "error": f"Failed to add chart: {str(e)}"
            }