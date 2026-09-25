"""
Visual inspection of generated presentations (render + vision-LLM review).

Pipeline (inspired by the render/inspect loop of document-generation agents,
implemented independently):
1. Render the .pptx to one PNG per slide: LibreOffice headless converts the
   deck to PDF, PyMuPDF rasterizes the pages. LibreOffice must be installed
   (`soffice` on PATH, or SOFFICE_PATH env var); the Dockerfile includes it.
2. Send the slide images to a configurable external vision LLM with a
   fidelity/error checklist, and parse a structured verdict.

The LLM endpoint speaks the OpenAI Responses API with image input
(Azure OpenAI included). Configuration via environment (see .env.example):
- VISION_LLM_ENDPOINT   full URL, e.g.
  https://<resource>.openai.azure.com/openai/responses?api-version=2025-04-01-preview
- VISION_LLM_API_KEY    sent as both api-key (Azure) and Authorization: Bearer
- VISION_LLM_MODEL      model / Azure deployment name (must accept images)
- VISION_LLM_BATCH_SLIDES slides per vision request (default 6); a deck is
  reviewed in batches, VISION_LLM_MAX_PARALLEL (default 4) at a time
- VISION_LLM_MAX_SLIDES cap on slides one whole-deck review covers (default 60)
- SOFFICE_PATH          LibreOffice binary if not "soffice" on PATH
"""
import base64
import contextvars
import io
import json
import logging
import os
import re
import shutil
import subprocess
import tempfile
import threading
import time
from pathlib import Path

import httpx
from pptx import Presentation

import deck_review
import fonts
from logging_utils import get_logger, flatten

logger = get_logger("visual_qa")


class VisualQAError(RuntimeError):
    pass


# ---- Rendering ----

def _soffice_binary():
    path = os.environ.get("SOFFICE_PATH") or shutil.which("soffice")
    if not path or not os.path.exists(path):
        raise VisualQAError(
            "LibreOffice is required for slide rendering but was not found. "
            "Install it (the server Docker image includes it) or set "
            "SOFFICE_PATH to the soffice binary."
        )
    return path


_slots = None
_slots_guard = threading.Lock()


def _conversion_slots():
    """Semaphore bounding concurrent LibreOffice processes.

    Read lazily rather than at import so the .env file loaded in main() is
    already in effect.
    """
    global _slots
    with _slots_guard:
        if _slots is None:
            raw = os.environ.get("PPT_MCP_MAX_CONCURRENT_CONVERSIONS",
                                 "").strip()
            limit = 2
            if raw:
                try:
                    limit = max(1, int(raw))
                except ValueError:
                    logger.warning("invalid_max_concurrent_conversions "
                                   "value=%r falling_back_to_default", raw)
            logger.debug("conversion_limit slots=%d", limit)
            _slots = threading.Semaphore(limit)
        return _slots


def convert_with_soffice(data: bytes, source_suffix: str, target: str,
                         timeout: float = 180.0) -> bytes:
    """Run one headless LibreOffice conversion and return the output bytes.

    `source_suffix` is the input file's extension (".pptx", ".ppt"); `target`
    is LibreOffice's --convert-to argument ("pdf", "pptx"). Every caller here
    goes through this: the isolated user profile is what stops concurrent
    conversions fighting over the shared profile lock, and getting that wrong
    fails intermittently under load rather than in testing.
    """
    with tempfile.TemporaryDirectory(prefix="pptx-convert-") as tmp:
        tmp = Path(tmp)
        src = tmp / f"deck{source_suffix}"
        src.write_bytes(data)
        profile = tmp / "lo-profile"
        home = tmp / "home"
        (home / ".cache").mkdir(parents=True, exist_ok=True)
        (home / ".config").mkdir(parents=True, exist_ok=True)
        cmd = [
            _soffice_binary(), "--headless", "--norestore",
            f"-env:UserInstallation=file://{profile}",
            "--convert-to", target, "--outdir", str(tmp), str(src),
        ]
        # dconf and fontconfig key their caches off $HOME, not off the
        # UserInstallation profile, so a deployment whose $HOME is unset,
        # read-only or owned by another uid makes soffice fail before it
        # loads the document. Pointing $HOME at the per-conversion temp dir
        # keeps that working without the deployment having to supply one.
        env = {**os.environ, "HOME": str(home),
               "XDG_CACHE_HOME": str(home / ".cache"),
               "XDG_CONFIG_HOME": str(home / ".config")}
        started = time.monotonic()
        # One conversion peaks around half a gigabyte, so an unbounded burst
        # of them is what exhausts a pod's memory limit rather than its CPU
        # (which PPT_MCP_MAX_CONCURRENT_TOOL_CALLS already bounds). Queue
        # here instead: waiting is not counted against `timeout`, which
        # measures the conversion itself.
        with _conversion_slots():
            proc = subprocess.run(cmd, capture_output=True, timeout=timeout,
                                  env=env)
        # --convert-to may take a filter suffix ("pdf:impress_pdf_Export");
        # the file it writes is named after the bare extension.
        out = tmp / f"deck.{target.split(':', 1)[0]}"
        if proc.returncode != 0 or not out.exists():
            logger.error("convert_failed stage=libreoffice target=%s "
                         "returncode=%d bytes=%d stderr=%s", target,
                         proc.returncode, len(data),
                         flatten(proc.stderr.decode(errors="replace")[-300:]))
            raise VisualQAError(
                f"LibreOffice failed to convert the presentation to {target}: "
                + proc.stderr.decode(errors="replace")[-500:]
            )
        result = out.read_bytes()
        logger.debug("convert_ok target=%s in_bytes=%d out_bytes=%d "
                     "duration_ms=%d", target, len(data), len(result),
                     int((time.monotonic() - started) * 1000))
        return result


def render_pptx_bytes_to_pdf(pptx_data: bytes) -> bytes:
    """Convert presentation bytes to PDF bytes."""
    return convert_with_soffice(pptx_data, ".pptx", "pdf")


def convert_legacy_ppt(data: bytes) -> bytes:
    """Convert a binary PowerPoint 97-2003 (.ppt) file to .pptx bytes.

    python-pptx reads only OOXML, so a legacy deck has to be converted before
    anything else in this server can touch it.
    """
    return convert_with_soffice(data, ".ppt", "pptx")


def render_pptx_bytes_to_pngs(pptx_data: bytes, dpi: int = 96,
                              max_slides: int = None,
                              slides: list = None) -> list:
    """Render presentation bytes to a list of PNG bytes.

    slides: 1-based slide numbers to rasterize (default: every slide).
    LibreOffice always converts the whole deck — there is no per-slide
    conversion — so a subset only skips rasterization and, more importantly,
    the vision call. max_slides still caps how many images come back.
    """
    import pymupdf

    started = time.monotonic()
    images = []
    with pymupdf.open(stream=render_pptx_bytes_to_pdf(pptx_data),
                      filetype="pdf") as doc:
        page_count = doc.page_count
        if slides is None:
            wanted = list(range(1, page_count + 1))
        else:
            wanted = [n for n in slides if 1 <= n <= page_count]
        if max_slides is not None:
            wanted = wanted[:max_slides]
        for number in wanted:
            pix = doc[number - 1].get_pixmap(dpi=dpi)
            images.append(pix.tobytes("png"))
    if slides is None and page_count > len(images):
        logger.warning("render_truncated rendered=%d slides=%d cap=%s "
                       "hint=raise_VISION_LLM_MAX_SLIDES",
                       len(images), page_count, max_slides)
    logger.debug("render_ok slides=%d of=%d dpi=%d duration_ms=%d",
                 len(images), page_count, dpi,
                 int((time.monotonic() - started) * 1000))
    return images


def normalize_slides(pres, slides):
    """Validate a caller-supplied 1-based slide selection against the deck.

    Returns a sorted, de-duplicated list, or None for "the whole deck".
    Raises ValueError naming the out-of-range numbers, so the tool layer can
    hand the agent an actionable message.
    """
    if slides is None:
        return None
    if isinstance(slides, int):
        slides = [slides]
    total = len(pres.slides)
    numbers, bad = [], []
    for value in slides:
        if isinstance(value, bool) or not isinstance(value, int):
            bad.append(value)
        elif 1 <= value <= total:
            numbers.append(value)
        else:
            bad.append(value)
    if bad:
        raise ValueError(
            f"Invalid slide number(s) {bad}: this presentation has {total} "
            f"slide(s), numbered 1-{total}."
        )
    if not numbers:
        return None
    return sorted(set(numbers))


# ---- Vision LLM client (OpenAI Responses API shape, Azure-compatible) ----

REVIEW_PROMPT = """You are a meticulous presentation QA reviewer. You are shown \
rendered slide images of a PowerPoint deck generated from a corporate template{ref_note}. \
The deck will be presented to an audience as it is: anything you miss ships.

Judge the rendered pixels. Beside the images you get an inventory of what each \
slide's file actually contains (title, text of every element, charts, tables); \
use it to catch content that exists but cannot be seen, and content that does \
not belong on its slide. Check EVERY slide for EVERY item below.

1. Empty, incomplete or wrong content.
   - A content slide that shows only its title, or whose body area is blank or \
nearly blank, is critical — it will be presented as a blank slide. (Title, \
section divider, statement, quote and closing slides are meant to be sparse.)
   - Content that does not match the slide title — e.g. cards, figures or text \
about a different topic, the typical result of content placed on the wrong \
slide — is critical. Say what the content is about and which title it would fit.
   - Placeholder prompts ("Click to add title/text"), sample or template text, \
"Lorem ipsum", "TBD", "XXX", "[insert …]", empty picture or chart frames.
   - Large unexplained empty regions, or content crammed into one corner while \
the rest of the slide stays empty.
2. Hidden, covered or stray content.
   - Text from the inventory that is not visible: covered by a shape, card or \
picture, placed off the slide, same colour as its background, or shrunk to \
nothing. Text showing through behind or between cards (letters peeking out \
from under shapes) is critical.
   - Elements stacked on top of each other, duplicated elements, leftover \
shapes or empty boxes with no purpose, template decoration covered by content.
3. Text placement and overlap, ANYWHERE text appears — not just in text boxes:
   - Text overflowing, clipped, or spilling outside its container or the slide edge.
   - Text overlapping other text, or sitting on top of shapes, images or lines in \
a way that makes either hard to read.
   - Charts and graphs: axis tick labels colliding with each other or truncated \
("..." or cut-off words), data labels overlapping their bars/slices/points or each \
other, a legend covering the plot area or running off the chart, an axis title \
squeezed or rotated into illegibility, series labels detached from what they label.
   - Tables: cell text wrapping into an unreadable stack, clipped by the cell or \
row height, columns too narrow for their content, headers not aligned with their \
columns, a table extending past the slide, empty rows or columns.
   - Diagrams, SmartArt and grouped shapes: labels wider than the node or box that \
holds them, text escaping a connector or arrow, node labels overlapping neighbouring \
nodes or connectors.
   - Footnotes and source lines clipped at the slide bottom, running into \
content, or wrapping over the template's footer.
4. Text size and legibility.
   - Text too small to read at presentation size (body text below roughly 10pt, \
chart and table text below roughly 8pt), or too low-contrast against what is \
behind it (light on light, dark on dark, text over a busy photo).
   - Text sized badly for the space it occupies: a heading or body block set so \
small that its box is mostly empty, or comparable elements on one slide set at \
visibly different sizes for no reason. Text should fill its container \
comfortably without crowding it or its neighbours — report both the starved and \
the overstuffed cases, but do not ask for larger text where growing it would \
eat the slide's white space.
   - Awkward line breaks: a single word orphaned on its own line in a title or \
card, a word or number split across lines.
5. Characters and formatting.
   - Missing-glyph boxes (□, ▯, ?), mojibake ("â€™", "Ã©"), literal markup or \
escapes ("**bold**", "\\\\n", "&amp;", "<b>").
   - Doubled bullets (a typed "•" or "-" after an automatic bullet), empty bullet \
lines, a leading blank line, numbering that restarts or skips.
   - Obvious typos, doubled words, sentences cut off mid-word.
6. Charts and data visuals.
   - A chart squashed into a strip, its plot area too small to read, or empty \
(no bars, lines or slices).
   - Default or meaningless labels: "Chart Title", "Series 1", "Metric", \
"Category 1", an axis title that says nothing.
   - A chart whose form does not fit its data (one bar, a pie with one slice, \
a single point on a line), or whose scale hides the differences it is meant to show.
   - KPI or stat cards whose numbers are cut, misaligned or visually unrelated \
to their labels.
7. Pictures and icons.
   - Distorted aspect ratio, pixelated or blurry, awkwardly cropped (cut-off \
heads or text), a broken-image placeholder.
   - Icons inconsistent in style, colour or size across a row, or unreadable \
against their background.
8. Layout, alignment and consistency.
   - Elements off the slide edge, crossing the template's header rule, logo or \
footer zone, or covering template artwork.
   - Comparable elements (cards, columns, icons, timeline steps) misaligned, \
unevenly spaced or of unequal size; an unbalanced composition.
   - Title position, size or colour differing from the other slides of the \
same layout.
9. Template and brand fidelity: consistent colors, fonts, logo placement, and \
layout usage matching the deck's own master style{ref_clause}; off-brand colours \
or fonts; a missing, duplicated or distorted logo.

Severity:
- "critical": the slide would be presented wrong — blank or near-blank content \
slide, content that belongs to another slide, hidden or covered content, \
unreadable text, a broken or empty chart/table/image.
- "major": a defect an audience would notice — clipped or overlapping text \
(unreadable or overlapping text is at least major), placeholder text, \
distorted images, misaligned comparable elements, default chart labels.
- "minor": polish only.

Report each problem separately, naming the element it affects (e.g. "chart on the \
right: x-axis labels overlap", or the element's shape_index from the \
inventory), and say in "suggested_fix" what change would resolve it (resize, \
reposition, shorten the text, smaller font, wider column, hide the legend, \
move element N to slide M, delete the covered text, bring to front). When the \
fix needs content only the author can supply, say exactly what is missing.

Respond with ONLY a JSON object, no markdown fence:
{{"passed": true|false, "issues": [{{"slide": <1-based number>, "severity": \
"critical"|"major"|"minor", "category": "empty"|"misplaced"|"hidden"|"overflow"|\
"overlap"|"legibility"|"characters"|"chart"|"table"|"image"|"layout"|"brand", \
"element": "...", "related_slides": [<other 1-based slides involved, if any>], \
"description": "...", "suggested_fix": "..."}}]}}
"passed" is true only when there are no critical or major issues."""


class VisionLLMConfigError(VisualQAError):
    pass


# Azure OpenAI (and DIAL Core's Azure upstream) require an api-version on
# every request. This preview version covers both the Responses API and
# chat completions with image input. Overridable with VISION_LLM_API_VERSION.
DEFAULT_API_VERSION = "2025-04-01-preview"


def _with_api_version(url: str) -> str:
    """Ensure the request URL carries an ?api-version= query parameter.

    Azure OpenAI rejects requests without it ("api-version is a required
    query parameter"), and DIAL Core passes the parameter through to its
    Azure upstream, so both providers need it. A version already present in
    the configured URL always wins; otherwise VISION_LLM_API_VERSION (or the
    default) is appended.
    """
    from urllib.parse import parse_qs, urlparse

    if "api-version" in parse_qs(urlparse(url).query):
        return url
    version = os.environ.get("VISION_LLM_API_VERSION", DEFAULT_API_VERSION).strip()
    if not version:
        return url
    return f"{url}{'&' if urlparse(url).query else '?'}api-version={version}"


def _resolve_provider() -> str:
    """Which backend serves the vision LLM:
    - "direct": VISION_LLM_ENDPOINT + VISION_LLM_API_KEY (OpenAI Responses
      API, Azure OpenAI included) — the default whenever an endpoint is set.
    - "dial": the model is a DIAL Core deployment, called at
      {DIAL_CORE_URL}/openai/deployments/{model}/chat/completions with DIAL
      credentials (caller headers first, DIAL_API_KEY fallback — the same
      resolution as file storage).
    VISION_LLM_PROVIDER=direct|dial overrides the inference."""
    value = os.environ.get("VISION_LLM_PROVIDER", "").lower()
    if value in ("direct", "azure", "openai"):
        return "direct"
    if value in ("dial", "dial-core", "dial_core"):
        return "dial"
    return "direct" if os.environ.get("VISION_LLM_ENDPOINT") else "dial"


class VisionLLM:
    def __init__(self):
        self.model = os.environ.get("VISION_LLM_MODEL")
        self.provider = _resolve_provider()
        self.endpoint = os.environ.get("VISION_LLM_ENDPOINT")
        self.api_key = os.environ.get("VISION_LLM_API_KEY")
        self.dial_url = os.environ.get("DIAL_CORE_URL")
        if not self.model:
            raise VisionLLMConfigError(
                "Visual inspection is not configured: set VISION_LLM_MODEL "
                "(see .env.example)."
            )
        if self.provider == "direct" and not (self.endpoint and self.api_key):
            raise VisionLLMConfigError(
                "Visual inspection (direct provider) needs VISION_LLM_ENDPOINT "
                "and VISION_LLM_API_KEY (see .env.example)."
            )
        if self.provider == "dial" and not self.dial_url:
            raise VisionLLMConfigError(
                "Visual inspection (dial provider) needs DIAL_CORE_URL so the "
                "model can be called as a DIAL deployment (see .env.example)."
            )

    def build_payload(self, images: list, prompt: str) -> dict:
        if self.provider == "dial":
            content = [{"type": "text", "text": prompt}]
            for png in images:
                b64 = base64.b64encode(png).decode()
                content.append({
                    "type": "image_url",
                    "image_url": {"url": f"data:image/png;base64,{b64}",
                                  "detail": "high"},
                })
            return {"messages": [{"role": "user", "content": content}]}
        content = [{"type": "input_text", "text": prompt}]
        for png in images:
            b64 = base64.b64encode(png).decode()
            content.append({
                "type": "input_image",
                "image_url": f"data:image/png;base64,{b64}",
                # "auto" lets the endpoint downscale a multi-image request
                # until small text, clipped glyphs and thin overlaps — most
                # of what the reviewer is there to find — are no longer
                # visible.
                "detail": "high",
            })
        return {"model": self.model,
                "input": [{"role": "user", "content": content}]}

    def _request_target(self):
        """(url, headers) for the configured provider."""
        if self.provider == "dial":
            from dial_client import DialConfigError, resolve_dial_auth_headers
            try:
                headers = resolve_dial_auth_headers()
            except DialConfigError as e:
                raise VisionLLMConfigError(str(e))
            headers["Content-Type"] = "application/json"
            url = (f"{self.dial_url.rstrip('/')}/openai/deployments/"
                   f"{self.model}/chat/completions")
            return _with_api_version(url), headers
        return _with_api_version(self.endpoint), {
            "api-key": self.api_key,                       # Azure OpenAI
            "Authorization": f"Bearer {self.api_key}",     # OpenAI-compatible
            "Content-Type": "application/json",
        }

    def extract_text(self, response_json: dict) -> str:
        """Pull the assistant text out of the provider's response."""
        if self.provider == "dial":
            try:
                content = response_json["choices"][0]["message"]["content"]
            except (KeyError, IndexError, TypeError):
                raise VisualQAError("Vision LLM returned no text output.")
            if isinstance(content, list):  # multimodal content parts
                content = "\n".join(
                    c.get("text", "") for c in content
                    if isinstance(c, dict) and c.get("type") == "text")
            if not isinstance(content, str) or not content:
                raise VisualQAError("Vision LLM returned no text output.")
            return content
        if isinstance(response_json.get("output_text"), str):
            return response_json["output_text"]
        parts = []
        for item in response_json.get("output", []):
            if item.get("type") == "message":
                for c in item.get("content", []):
                    if c.get("type") == "output_text":
                        parts.append(c.get("text", ""))
        if not parts:
            raise VisualQAError("Vision LLM returned no text output.")
        return "\n".join(parts)

    @staticmethod
    def parse_verdict(text: str) -> dict:
        """Parse the reviewer's JSON verdict, tolerating a stray code fence."""
        candidate = text.strip()
        m = re.search(r"\{.*\}", candidate, re.S)
        if m:
            candidate = m.group(0)
        try:
            verdict = json.loads(candidate)
            if isinstance(verdict, dict) and "passed" in verdict:
                verdict.setdefault("issues", [])
                return verdict
        except json.JSONDecodeError:
            pass
        logger.warning("vision_verdict_unparseable chars=%d preview=%s",
                       len(text), flatten(text[:200]))
        return {"passed": None, "issues": [],
                "raw_review": text,
                "note": "Reviewer response was not valid JSON; see raw_review."}

    def ask(self, images: list, prompt: str, timeout: float = 300.0) -> str:
        """Send prompt + images, return the model's raw text answer."""
        url, headers = self._request_target()
        payload = self.build_payload(images, prompt)
        logger.debug("vision_request provider=%s model=%s images=%d "
                     "prompt_chars=%d timeout_s=%.0f",
                     self.provider, self.model, len(images), len(prompt), timeout)
        started = time.monotonic()
        r = httpx.post(url, headers=headers, json=payload, timeout=timeout)
        duration_ms = int((time.monotonic() - started) * 1000)
        if r.status_code != 200:
            detail = r.text[:300]
            logger.error("vision_request_failed provider=%s model=%s status=%d "
                         "duration_ms=%d detail=%s", self.provider, self.model,
                         r.status_code, duration_ms, flatten(detail))
            if "api-version" in detail:
                detail += (" — set VISION_LLM_API_VERSION to a version your "
                           "endpoint accepts (or put ?api-version=... in "
                           "VISION_LLM_ENDPOINT).")
            raise VisualQAError(
                f"Vision LLM request failed with HTTP {r.status_code}: {detail}"
            )
        text = self.extract_text(r.json())
        logger.debug("vision_response provider=%s model=%s duration_ms=%d "
                     "chars=%d", self.provider, self.model, duration_ms, len(text))
        return text

    def review(self, images: list, prompt: str, timeout: float = 300.0) -> dict:
        return self.parse_verdict(self.ask(images, prompt, timeout))

    def ask_json(self, images: list, prompt: str, timeout: float = 300.0) -> dict:
        """Like ask(), parsed as a JSON object ({} when unparseable)."""
        text = self.ask(images, prompt, timeout)
        candidate = text.strip()
        m = re.search(r"\{.*\}", candidate, re.S)
        if m:
            candidate = m.group(0)
        try:
            data = json.loads(candidate)
            return data if isinstance(data, dict) else {}
        except json.JSONDecodeError:
            logger.warning("vision_json_unparseable chars=%d preview=%s",
                           len(text), flatten(text[:200]))
            return {}


def vision_configured() -> bool:
    """Whether a vision model can be reached at all — either directly
    (VISION_LLM_ENDPOINT + VISION_LLM_API_KEY) or as a DIAL Core deployment
    (DIAL_CORE_URL).

    Deliberately separate from enforcement_enabled(): VISUAL_QA_ENFORCE turns
    off *slide* inspection and repair, and a feature that merely uses the model
    for its own check — the icon review — should follow whether the model
    exists, not that switch."""
    if not os.environ.get("VISION_LLM_MODEL"):
        return False
    if _resolve_provider() == "direct":
        return bool(os.environ.get("VISION_LLM_ENDPOINT")
                    and os.environ.get("VISION_LLM_API_KEY"))
    return bool(os.environ.get("DIAL_CORE_URL"))


def enforcement_enabled() -> bool:
    """Visual QA is available when the vision LLM is configured, unless
    disabled with VISUAL_QA_ENFORCE=false.

    Gates registration of the inspect/repair tools; also required for the
    optional export gate (see export_gate_enabled)."""
    return vision_configured() and os.environ.get(
        "VISUAL_QA_ENFORCE", "true").lower() != "false"


def export_gate_enabled() -> bool:
    """Whether export/save should run the inspect-repair loop themselves.

    Off by default: QA is driven by the orchestrator through the
    visual_inspect_slides / visual_repair_slides tools, which it can call as
    often as it likes on whichever slides it just built. Operators who want
    the old always-on gate — a deck can never leave the server uninspected —
    set VISUAL_QA_EXPORT_GATE=true.
    """
    return (enforcement_enabled()
            and os.environ.get("VISUAL_QA_EXPORT_GATE", "false").lower() == "true")


def unresolved_policy() -> str:
    """What to do when the internal repair loop exhausts its iterations
    without a pass: "report" (default — the export fails with the unresolved
    issue list) or "export_as_is" (ship the best-effort deck).
    VISUAL_QA_ON_UNRESOLVED accepts report/block and export_as_is/export."""
    value = os.environ.get("VISUAL_QA_ON_UNRESOLVED", "report").lower()
    return "export_as_is" if value in ("export", "export_as_is") else "report"


def fail_open_on_error() -> bool:
    """VISUAL_QA_ON_ERROR=allow lets exports through when inspection itself
    fails (renderer missing, endpoint down). Default is to block."""
    return os.environ.get("VISUAL_QA_ON_ERROR", "block").lower() == "allow"


def _prune_unused_layouts(pres):
    """Remove every slide layout, and every slide master, that no slide in
    `pres` uses. Mutates `pres` — only ever call it on a throwaway copy.

    A corporate template can carry hundreds of layouts across several
    masters, and LibreOffice imports every one of them as a master page
    before it renders a single slide. That import, not the slides, is what
    dominates conversion time: a 1-slide render of a 291-layout template
    takes ~15s, the same slide with its one layout ~1s. The rendered pixels
    are identical, because a slide only ever draws its own layout and master.
    """
    if not len(pres.slides):
        return
    used_layouts = {s.slide_layout.part.partname for s in pres.slides}
    used_masters = {s.slide_layout.slide_master.part.partname
                    for s in pres.slides}
    for master in list(pres.slide_masters):
        layout_ids = master._element.get_or_add_sldLayoutIdLst()
        for layout_id in list(layout_ids):
            rid = layout_id.rId
            if master.part.related_part(rid).partname not in used_layouts:
                # Remove the reference first: drop_rel only drops a
                # relationship nothing in the XML still points at.
                layout_ids.remove(layout_id)
                master.part.drop_rel(rid)
    master_ids = pres.part._element.get_or_add_sldMasterIdLst()
    for master_id in list(master_ids):
        rid = master_id.rId
        if pres.part.related_part(rid).partname not in used_masters:
            master_ids.remove(master_id)
            pres.part.drop_rel(rid)


def _subset_deck_bytes(pres, slides=None):
    """Serialize a render-only copy of `pres`: just the given 1-based slides
    (None keeps them all), kept in their original deck order (not the order
    of `slides`), with every unused layout and master stripped.

    LibreOffice always converts the whole file it is given — slides, layouts
    and masters alike — so trimming here, before the file ever reaches
    soffice, is what actually cuts render time. Operates on a freshly
    reopened copy so the caller's live `pres` (and its bound `shapes`/spTree
    state, see CLAUDE.md) is never touched.
    """
    import io
    from pptx import Presentation
    from utils import delete_slide

    buf = io.BytesIO()
    pres.save(buf)
    buf.seek(0)
    subset = Presentation(buf)
    if slides:
        keep = set(slides)
        for index in range(len(subset.slides) - 1, -1, -1):
            if (index + 1) not in keep:
                delete_slide(subset, index)
    try:
        _prune_unused_layouts(subset)
    except Exception as e:  # never let an optimization break a render
        logger.warning("layout_prune_failed error=%s falling_back=unpruned",
                       flatten(str(e)))
        buf.seek(0)
        return _subset_deck_bytes_unpruned(buf.getvalue(), slides)
    out = io.BytesIO()
    subset.save(out)
    return out.getvalue()


def _subset_deck_bytes_unpruned(data, slides):
    """The pre-pruning behaviour, kept as the fallback path."""
    import io
    from pptx import Presentation
    from utils import delete_slide

    subset = Presentation(io.BytesIO(data))
    if slides:
        keep = set(slides)
        for index in range(len(subset.slides) - 1, -1, -1):
            if (index + 1) not in keep:
                delete_slide(subset, index)
    out = io.BytesIO()
    subset.save(out)
    return out.getvalue()


def _render_deck(pres, max_slides=None, slides=None):
    # slides is always pre-sorted (normalize_slides / the repair loop's own
    # review scope), so the subset deck's page order already matches what
    # callers expect back — no slides= filtering needed downstream. A
    # whole-deck render goes through the same pruned copy, so it too skips
    # the template's unused layouts.
    return render_pptx_bytes_to_pngs(_subset_deck_bytes(pres, slides),
                                     max_slides=None if slides else max_slides)


def _env_int(name, default, minimum=1):
    try:
        return max(minimum, int(os.environ.get(name, default)))
    except ValueError:
        logger.warning("config_invalid var=%s value=%s using=%s", name,
                       os.environ.get(name), default)
        return default


def _slide_cap():
    """Most slides one whole-deck review covers. Slides beyond it are reported
    back as not reviewed — never silently treated as passed."""
    return _env_int("VISION_LLM_MAX_SLIDES", 60)


def _batch_size():
    """Slides per vision request. Fifteen images in one request is where the
    reviewer started missing blank slides outright: each image gets a smaller
    share of the model's attention (and of the endpoint's image budget)."""
    return _env_int("VISION_LLM_BATCH_SLIDES", 6)


def _max_parallel():
    return _env_int("VISION_LLM_MAX_PARALLEL", 4)


def _initial_scope(pres, slides):
    """-> (review_scope, not_reviewed). review_scope is None for "the whole
    deck, rendered as one" and a sorted list otherwise; not_reviewed lists
    the slides a whole-deck review had to leave out because of the cap."""
    if slides:
        return slides, []
    total, cap = len(pres.slides), _slide_cap()
    if total <= cap:
        return None, []
    logger.warning("review_truncated slides=%d cap=%d "
                   "hint=raise_VISION_LLM_MAX_SLIDES", total, cap)
    return list(range(1, cap + 1)), list(range(cap + 1, total + 1))


def _coherence_applies(pres, slides):
    """The deck-level story review needs the whole deck — an agenda cannot be
    checked against the two slides a scoped call names — and something to
    cross-check."""
    return (slides is None and len(pres.slides) >= 3
            and deck_review.coherence_enabled())


def _valid_issue(issue, allowed):
    return isinstance(issue, dict) and issue.get("slide") in allowed


def _review(llm, pres, images, image_slides, focus, risky_fonts,
            reference_images=(), coherence=False):
    """Review rendered slides in parallel batches, plus (optionally) the
    deck-level coherence review, and merge the results.

    Returns (visual, coherence_verdict): visual is {"passed", "issues",
    "slides_reviewed"} with passed None when any batch answer could not be
    parsed; coherence_verdict is None when not run or when it failed (a
    failed story review degrades to the visual one, it does not fail QA).
    """
    from concurrent.futures import ThreadPoolExecutor

    # The inventory is built here, on the calling thread: the review threads
    # only do HTTP and never touch python-pptx objects.
    inventory = {o["slide"]: o
                 for o in deck_review.slide_outline(pres, image_slides)}
    outline = deck_review.slide_outline(pres) if coherence else None
    size = _batch_size()
    batches = [(images[i:i + size], image_slides[i:i + size])
               for i in range(0, len(images), size)]
    reference_images = list(reference_images)

    def review_batch(batch):
        batch_images, numbers = batch
        prompt = review_prompt(
            bool(reference_images), focus, numbers, risky_fonts,
            inventory=[inventory[n] for n in numbers if n in inventory],
            image_offset=len(reference_images))
        return numbers, llm.review(reference_images + batch_images, prompt)

    def submit(pool, fn, *args):
        # Pool threads start with an empty context, and the caller's DIAL
        # credentials live in the MCP SDK's request contextvar: without a
        # copy, the DIAL provider sees no request and no Api-Key. Each task
        # needs its own copy — one Context cannot be entered by two threads.
        return pool.submit(contextvars.copy_context().run, fn, *args)

    workers = max(1, min(_max_parallel(), len(batches) + bool(coherence)))
    with ThreadPoolExecutor(max_workers=workers) as pool:
        coherence_future = submit(
            pool, llm.review, [], deck_review.coherence_prompt(outline)) \
            if coherence else None
        batch_futures = [submit(pool, review_batch, b) for b in batches]
        batch_results = [f.result() for f in batch_futures]

    issues, unparseable = [], None
    for numbers, verdict in batch_results:
        if verdict.get("passed") is None and unparseable is None:
            unparseable = verdict
        for issue in verdict.get("issues", []):
            if _valid_issue(issue, numbers):
                issue.setdefault("check", "visual")
                issues.append(issue)
            else:
                logger.debug("review_issue_dropped slide=%s batch=%s",
                             issue.get("slide") if isinstance(issue, dict)
                             else "?", ",".join(map(str, numbers)))
    visual = {"passed": None if unparseable else not blocking_issues(issues),
              "issues": issues,
              "slides_reviewed": list(image_slides)}
    if unparseable:
        for key in ("raw_review", "note"):
            if key in unparseable:
                visual[key] = unparseable[key]

    coherence_verdict = None
    if coherence_future is not None:
        try:
            coherence_verdict = coherence_future.result()
            valid = range(1, len(pres.slides) + 1)
            coherence_verdict["issues"] = [
                dict(i, check="coherence")
                for i in coherence_verdict.get("issues", [])
                if _valid_issue(i, valid)]
            logger.info("coherence_done slides=%d passed=%s issues=%d",
                        len(pres.slides), coherence_verdict.get("passed"),
                        len(coherence_verdict["issues"]))
        except Exception as e:
            logger.warning("coherence_failed error=%s", flatten(str(e)))
            coherence_verdict = None
    return visual, coherence_verdict


def inspect_presentation(pres, reference_pres=None, focus: str = None,
                         slides: list = None) -> dict:
    """Render a python-pptx Presentation (and optional reference) and return
    the reviewers' verdict.

    slides: 1-based slide numbers to review; None reviews the whole deck (up
    to VISION_LLM_MAX_SLIDES, the rest reported in "slides_not_reviewed") and
    also runs the deck-level coherence review. Issue slide numbers are always
    absolute deck positions. "passed" is decided here from the severities —
    never taken from the model's own flag.
    Raises VisualQAError on infrastructure failure (renderer/LLM).
    """
    llm = VisionLLM()
    scope, not_reviewed = _initial_scope(pres, slides)
    deck_images = _render_deck(pres, None, scope)
    image_slides = scope or list(range(1, len(deck_images) + 1))
    # One batch's worth of reference pages is enough to show the brand.
    ref_images = _render_deck(reference_pres, None, None)[:_batch_size()] \
        if reference_pres is not None else []

    visual, coherence = _review(
        llm, pres, deck_images, image_slides, focus,
        fonts.unreliable_fonts_in(pres), ref_images,
        coherence=_coherence_applies(pres, slides))
    issues = visual["issues"] + (coherence["issues"] if coherence else [])
    verdict = {
        "passed": None if visual["passed"] is None
        else not blocking_issues(issues),
        "issues": issues,
        "slides_reviewed": len(image_slides),
        "checks": ["visual"] + (["coherence"] if coherence else []),
    }
    if not_reviewed:
        verdict["slides_not_reviewed"] = not_reviewed
    for key in ("raw_review", "note"):
        if key in visual:
            verdict[key] = visual[key]
    logger.info("inspection_done slides=%d scope=%s reference=%s passed=%s "
                "issues=%d", len(deck_images),
                ",".join(map(str, slides)) if slides else "deck",
                bool(ref_images), verdict["passed"], len(issues))
    return verdict


# Only these make a slide fail. "minor" findings are reported back but never
# repaired: a stateless reviewer always finds another minor nit, and chasing
# them is what turned one-round fixes into ten-round loops. An issue with no
# (or an unknown) severity counts as blocking, so nothing slips through.
NON_BLOCKING_SEVERITIES = {"minor"}
# Cap on minor findings echoed back: they are for the record, not for action.
MAX_MINOR_REPORTED = 20


def blocking_issues(issues):
    return [i for i in issues
            if str(i.get("severity", "")).lower() not in NON_BLOCKING_SEVERITIES]


def inspect_and_repair(pres, slides: list = None, focus: str = None,
                       max_iterations: int = None) -> dict:
    """Inspect/repair loop: review the selected slides (and, for the whole
    deck, the deck's story — see deck_review.py); on failure, repair in place
    via LLM-planned whitelisted operations (visual_fix.py) and review again,
    up to VISUAL_QA_MAX_ITERATIONS (default 3) reviews. Stops early once a
    repair round fails to reduce the blocking-issue count: the operations
    either fixed it or they cannot. A round that makes a slide worse is
    rolled back on that slide (_judge_round), so the deck left behind holds
    the best version of each slide seen, not the last attempt.

    slides: 1-based slide numbers to work on; None means the whole deck.
    Repairs are confined to the reviewed slides — issues reported against
    other slides are ignored, so a caller iterating slide by slide never has
    the model rewrite a slide it did not ask about.

    Returns {"passed", "iterations", "repair_rounds", "issues",
    "minor_issues", "slides_reviewed", "checks"} plus "action_required" (what
    only the deck's author can fix) and "slides_not_reviewed" when non-empty.
    "passed" is our rule — no blocking issue left anywhere in scope — never
    the model's own flag, which it happily sets beside a critical finding.
    Raises VisualQAError on infrastructure failure (renderer/LLM).
    """
    import visual_fix

    llm = VisionLLM()
    if max_iterations is None:
        max_iterations = int(os.environ.get("VISUAL_QA_MAX_ITERATIONS", "3"))
    # The budget counts inspections, so 1 would inspect and return without
    # ever repairing — which is what an orchestrator asking for "one quick
    # round" means by 1 and never gets. Two is the smallest budget that
    # repairs: inspect, repair, re-inspect.
    max_iterations = max(2, max_iterations)

    # Constant for the whole loop: repairs never change which fonts the deck
    # names, and re-scanning per round would only cost time.
    risky_fonts = fonts.unreliable_fonts_in(pres)
    initial_scope, not_reviewed = _initial_scope(pres, slides)
    use_coherence = _coherence_applies(pres, slides)
    checks = ["visual"] + (["coherence"] if use_coherence else [])
    repair_rounds, author_actions = [], []
    visual = {}
    loop_started = time.monotonic()
    logger.info("qa_loop_start scope=%s coherence=%s max_iterations=%d",
                ",".join(map(str, slides)) if slides else "deck",
                use_coherence, max_iterations)
    # Round 1 reviews everything in scope. Later rounds re-render and
    # re-review only the slides the previous round actually changed — an
    # untouched slide cannot have changed, and re-reviewing it just invites
    # a stateless reviewer to find a fresh nit on a slide that already
    # passed. Blocking visual issues on slides no operation reached are
    # carried forward as unresolved instead of being re-inspected. The
    # coherence review is one cheap text call and re-runs every round:
    # a move or an agenda rewrite changes the story on several slides.
    review_scope = initial_scope
    # The latest accepted verdict per slide. A slide keeps its issues until
    # it is reviewed again, which is how a blocking issue on a slide no
    # operation reached stays unresolved without being re-inspected.
    images_by_slide, visual_by_slide, coherence_issues = {}, {}, []
    reviewed = set()
    # The round whose outcome the next review judges: the deck as it was
    # before that round, the operations it applied and the per-slide state
    # to fall back on. See _judge_round.
    pending = None
    for iteration in range(1, max_iterations + 1):
        round_started = time.monotonic()
        deck_images = _render_deck(pres, None, review_scope)
        # Absolute slide number of each image, so issues and repairs address
        # deck positions even when only a subset was rendered.
        image_slides = review_scope or list(range(1, len(deck_images) + 1))
        visual, coherence = _review(llm, pres, deck_images, image_slides,
                                    focus, risky_fonts,
                                    coherence=use_coherence)
        new_visual = {n: [] for n in image_slides}
        for issue in visual["issues"]:
            if not slides or issue.get("slide") in slides:
                new_visual.setdefault(issue.get("slide"), []).append(issue)
        new_state = {"visual": new_visual,
                     "images": dict(zip(image_slides, deck_images)),
                     "coherence": coherence["issues"] if coherence else []}
        reverted = []
        if pending:
            reverted, stop = _judge_round(pres, pending, new_state, slides)
            if stop:
                # A reorder made the deck worse and was undone as a whole:
                # nothing reviewed this round describes the deck any more.
                visual_by_slide = pending["visual"]
                images_by_slide = pending["images"]
                coherence_issues = pending["coherence"]
                break
        visual_by_slide.update(new_state["visual"])
        images_by_slide.update(new_state["images"])
        coherence_issues = new_state["coherence"]
        reviewed.update(image_slides)

        found = _all_issues(visual_by_slide, coherence_issues)
        issues = blocking_issues(found)
        # What this round can act on: slides it just reviewed (and did not
        # just roll back — that fix was tried) plus the deck's story.
        actionable = [i for i in issues if i.get("check") == "coherence"
                      or (i.get("slide") in image_slides
                          and i.get("slide") not in reverted)]
        logger.info("qa_round iteration=%d/%d slides=%d issues=%d blocking=%d "
                    "actionable=%d reverted=%d coherence_issues=%s "
                    "duration_ms=%d", iteration, max_iterations,
                    len(deck_images), len(found), len(issues),
                    len(actionable), len(reverted),
                    len(coherence["issues"]) if coherence else "-",
                    int((time.monotonic() - round_started) * 1000))
        if logger.isEnabledFor(logging.DEBUG):
            for issue in found:
                logger.debug("qa_issue iteration=%d check=%s slide=%s "
                             "severity=%s description=%s", iteration,
                             issue.get("check"), issue.get("slide"),
                             issue.get("severity"),
                             flatten(str(issue.get("description", ""))[:200]))
        if not issues and visual["passed"] is not None:
            logger.info("qa_loop_passed iterations=%d repair_rounds=%d "
                        "duration_ms=%d", iteration, len(repair_rounds),
                        int((time.monotonic() - loop_started) * 1000))
            return _outcome(True, iteration, repair_rounds, [], found,
                            reviewed, checks, not_reviewed, author_actions)
        if repair_rounds and len(issues) >= repair_rounds[-1]["issues_found"]:
            # The last repair round left as many blocking issues as before:
            # more rounds of the same operations will not converge. Hand it
            # back instead of looping.
            logger.warning("qa_loop_stop reason=no_improvement iteration=%d "
                           "blocking=%d", iteration, len(issues))
            break
        if iteration == max_iterations or not actionable:
            # Out of budget, or nothing actionable (e.g. unparseable review,
            # or only unreached issues left)
            logger.warning("qa_loop_stop reason=%s iteration=%d",
                           "budget_exhausted" if iteration == max_iterations
                           else "no_actionable_issues", iteration)
            break
        author_actions = []
        plan = visual_fix.plan_repairs(
            llm, actionable, pres, list(images_by_slide.values()),
            list(images_by_slide.keys()), author_actions=author_actions)
        before = _deck_bytes(pres)
        result = visual_fix.apply_repairs(pres, plan, allowed_slides=slides)
        round_report = {
            "iteration": iteration,
            "issues_found": len(issues),
            "operations_applied": len(result["applied"]),
            "operations_skipped": len(result["skipped"]),
        }
        if result["applied"]:
            round_report["changes"] = visual_fix.describe_changes(
                result["applied"])
        if result["skipped"]:
            # Why nothing changed matters more than that nothing changed:
            # "bad shape_index" means the fix targets something the repair
            # engine cannot reach (a layout/master placeholder, say), which
            # no amount of re-running will improve.
            round_report["skipped_reasons"] = visual_fix.skip_reason_summary(
                result["skipped"])
        repair_rounds.append(round_report)
        if not result["applied"]:
            logger.warning("qa_loop_stop reason=no_repair_progress iteration=%d "
                           "operations_planned=%d operations_skipped=%d",
                           iteration, len(plan), len(result["skipped"]))
            break  # no progress is possible; stop burning inspections
        pending = {"deck": before, "applied": result["applied"],
                   "reordered": bool(result.get("slides_reordered")),
                   "touched": visual_fix.touched_slides(result["applied"]),
                   "report": round_report,
                   "visual": dict(visual_by_slide),
                   "images": dict(images_by_slide),
                   "coherence": list(coherence_issues)}
        if pending["reordered"]:
            # Every slide number just changed: nothing rendered or reviewed
            # still points at the right slide, so start over on the whole scope.
            images_by_slide, visual_by_slide = {}, {}
            review_scope = initial_scope
        else:
            review_scope = pending["touched"]

    found = _all_issues(visual_by_slide, coherence_issues)
    remaining = blocking_issues(found)
    out = _outcome(False, len(repair_rounds) + 1, repair_rounds, remaining,
                   found, reviewed, checks, not_reviewed, author_actions)
    if repair_rounds and not repair_rounds[-1]["operations_applied"] \
            and not repair_rounds[-1]["operations_skipped"]:
        out["repair_note"] = (
            "The repair planner found no operation that can fix what is "
            "left — typically content that was never built. Do what "
            "action_required says with the editing tools, then call "
            "visual_repair_slides on those slides. Repeating this call as it "
            "is will not help.")
    elif repair_rounds and not repair_rounds[-1]["operations_applied"]:
        # Tell the agent what a zero-applied round means, so it stops the
        # deck rather than re-running an identical call.
        out["repair_note"] = (
            "The last round changed nothing: every planned operation was "
            "rejected (" + ", ".join(
                f"{reason} x{count}" for reason, count
                in repair_rounds[-1].get("skipped_reasons", {}).items())
            + "). Repairs are working; these issues are just outside what "
              "they can reach — 'bad shape_index' usually means the target "
              "belongs to the slide layout or master rather than the slide. "
              "Fix the content yourself with the editing tools, or report "
              "the issue to the user. Repeating this call will not help."
        )
    logger.warning("qa_loop_failed iterations=%d repair_rounds=%d "
                   "unresolved_issues=%d duration_ms=%d",
                   out["iterations"], len(repair_rounds), len(remaining),
                   int((time.monotonic() - loop_started) * 1000))
    for key in ("raw_review", "note"):
        if key in visual:
            out[key] = visual[key]
    return out


# How much a blocking issue weighs when a round is judged slide by slide: a
# round that trades one major issue for a critical one made the slide worse.
_SEVERITY_WEIGHT = {"critical": 3}


def _all_issues(visual_by_slide, coherence_issues):
    return [i for n in sorted(visual_by_slide, key=str)
            for i in visual_by_slide[n]] + list(coherence_issues)


def _slide_scores(issues):
    scores = {}
    for issue in blocking_issues(issues):
        weight = _SEVERITY_WEIGHT.get(str(issue.get("severity", "")).lower(), 2)
        scores[issue.get("slide")] = scores.get(issue.get("slide"), 0) + weight
    return scores


def _deck_bytes(pres):
    buf = io.BytesIO()
    pres.save(buf)
    return buf.getvalue()


def _restore_deck(pres, data):
    """Put pres back to the state saved in data, in place. The store and
    every caller hold this object, so it is refilled rather than replaced:
    a python-pptx Presentation is a proxy whose whole state (element, part,
    lazily cached collections) lives in its instance __dict__."""
    fresh = Presentation(io.BytesIO(data))
    pres.__dict__.clear()
    pres.__dict__.update(fresh.__dict__)


def _judge_round(pres, pending, new_state, slides):
    """Keep what the last repair round improved and undo what it made worse.

    The reviewer judges each slide the round touched against that slide's
    verdict before the round. A planner that makes room for a caption by
    crushing the chart beside it scores worse, and without this the loop
    would hand back the crushed chart — the last round's deck is otherwise
    what the caller gets, whether or not it was an improvement.

    A worse slide is rolled back by restoring the pre-round deck and
    replaying the round's other operations, which is exact: the operations
    address the pre-round deck and are deterministic. Operations reaching
    a rolled-back slide go too, along with every other slide they touch (a
    move has two ends). Rolled-back slides take their earlier verdict and
    image back in new_state. A reorder renumbers every slide, so it is
    judged on the whole deck and undone whole.

    Returns (reverted slide numbers, stop): stop is True when the whole round
    was undone and the loop should end on the earlier state."""
    import visual_fix

    before = _slide_scores(_all_issues(pending["visual"], pending["coherence"]))
    after = _slide_scores(_all_issues(new_state["visual"],
                                      new_state["coherence"]))
    report = pending["report"]
    if pending["reordered"]:
        if sum(after.values()) <= sum(before.values()):
            return [], False
        _restore_deck(pres, pending["deck"])
        report.update(operations_applied=0,
                      operations_reverted=len(pending["applied"]),
                      reverted_slides="all", changes=[])
        logger.warning("qa_round_reverted iteration=%d scope=deck "
                       "operations=%d", report["iteration"],
                       len(pending["applied"]))
        return [], True

    reverted = {n for n in pending["touched"]
                if after.get(n, 0) > before.get(n, 0)}
    if not reverted:
        return [], False
    applied = pending["applied"]
    while True:
        dropped = [op for op in applied
                   if reverted & set(visual_fix.touched_slides([op]))]
        grown = reverted.union(*(visual_fix.touched_slides([op])
                                 for op in dropped))
        if grown == reverted:
            break
        reverted = grown
    dropped_ids = {id(op) for op in dropped}
    kept = [op for op in applied if id(op) not in dropped_ids]
    _restore_deck(pres, pending["deck"])
    if kept:
        visual_fix.apply_repairs(pres, kept, allowed_slides=slides)
    for n in reverted:
        new_state["visual"][n] = pending["visual"].get(n, [])
        if n in pending["images"]:
            new_state["images"][n] = pending["images"][n]
        else:
            new_state["images"].pop(n, None)
    new_state["coherence"] = (
        [i for i in new_state["coherence"] if i.get("slide") not in reverted]
        + [i for i in pending["coherence"] if i.get("slide") in reverted])
    report.update(operations_applied=len(kept),
                  operations_reverted=len(dropped),
                  reverted_slides=sorted(reverted),
                  changes=visual_fix.describe_changes(kept))
    logger.warning("qa_round_reverted iteration=%d slides=%s operations=%d "
                   "kept=%d", report["iteration"],
                   ",".join(map(str, sorted(reverted))), len(dropped),
                   len(kept))
    return sorted(reverted), False


def _outcome(passed, iterations, repair_rounds, remaining, found, reviewed,
             checks, not_reviewed, author_actions):
    minor = [i for i in found if i not in blocking_issues(found)]
    out = {"passed": passed,
           "iterations": iterations,
           "repair_rounds": repair_rounds,
           "issues": remaining,
           "minor_issues": minor[:MAX_MINOR_REPORTED],
           "slides_reviewed": len(reviewed),
           "checks": checks}
    if not passed and remaining:
        # What the repair engine cannot do by construction — write missing
        # content, invent a figure — goes back to the agent as instructions,
        # in the planner's words where it gave any.
        # A remaining issue on a slide the planner said nothing about still
        # needs an instruction: fall back to the reviewer's suggested fix.
        covered = {a.get("slide") for a in author_actions}
        actions = list(author_actions)
        for issue in remaining:
            if issue.get("slide") not in covered:
                covered.add(issue.get("slide"))
                actions.append({"slide": issue.get("slide"),
                                "action": issue.get("suggested_fix")
                                or issue.get("description")})
        out["action_required"] = actions
    if not_reviewed:
        out["slides_not_reviewed"] = not_reviewed
        out["review_note"] = (
            f"Slides {not_reviewed[0]}-{not_reviewed[-1]} were beyond "
            f"VISION_LLM_MAX_SLIDES and were not reviewed; run "
            f"visual_repair_slides with slides=[...] on them.")
    return out


def _inventory_text(entry):
    """One slide's inventory line for the vision reviewer."""
    parts = [f"slide {entry['slide']} (layout \"{entry['layout']}\")"]
    parts.append(f"title: {entry['title']!r}" if entry.get("title")
                 else "title: none")
    elements = []
    for element in entry["elements"]:
        if element.get("role") == "title":
            continue
        desc = f"#{element['shape_index']} {element['kind']}"
        if element.get("text"):
            desc += f": {element['text'][:160]!r}"
        elif element.get("empty"):
            desc += " (empty)"
        if element.get("chart"):
            desc += f" {json.dumps(element['chart'], ensure_ascii=False)[:160]}"
        if element.get("drawn_over_by"):
            desc += f" [drawn over by #{element['drawn_over_by']}]"
        elements.append(desc)
    parts.append("elements: " + ("; ".join(elements) if elements
                                 else "none besides the title"))
    return " — ".join(parts)


def review_prompt(has_reference: bool, focus: str = None,
                  slides: list = None, risky_fonts=None, inventory=None,
                  image_offset: int = 0) -> str:
    prompt = REVIEW_PROMPT.format(
        ref_note=(". The FIRST images are the reference template's slides; "
                  "the deck under review follows" if has_reference else ""),
        ref_clause=(" and matching the reference template images"
                    if has_reference else ""),
    )
    if slides:
        mapping = ", ".join(f"image {i} = slide {n}"
                            for i, n in enumerate(slides,
                                                  start=image_offset + 1))
        prompt += (
            f"\nYou are shown only part of a larger deck: {mapping}. Report "
            "every issue with the slide number given here, not the image "
            "position, and judge each slide on its own merits."
        )
    if inventory:
        prompt += ("\n\nWhat each slide's file contains (shape_index, kind, "
                   "text), to compare with what the image shows:\n"
                   + "\n".join(_inventory_text(e) for e in inventory))
    prompt += fonts.qa_font_caveat(risky_fonts)
    if focus:
        prompt += (f"\nAdditional focus requested by the caller (on top of, "
                   f"never instead of, the full checklist): {focus}")
    return prompt
