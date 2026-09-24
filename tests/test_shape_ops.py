"""Tests for the shape-level operations (utils/slide_utils.py) and the
manage_shape tool that exposes them."""
import io
import os
import sys
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches

import deck_validation
import utils as ppt_utils
from state import PresentationStore
from tools.slide_tools import register_slide_tools

PNG = (b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01\x00\x00\x00\x01"
       b"\x08\x02\x00\x00\x00\x90wS\xde\x00\x00\x00\x0cIDATx\x9cc\xf8\xcf"
       b"\xc0\x00\x00\x03\x01\x01\x00\xc9\xfe\x92\xef\x00\x00\x00\x00IEND"
       b"\xaeB`\x82")


def blank_deck(slides=2):
    pres = Presentation()
    for _ in range(slides):
        pres.slides.add_slide(pres.slide_layouts[6])
    return pres


def reopen(pres):
    buf = io.BytesIO()
    pres.save(buf)
    buf.seek(0)
    return Presentation(buf)


def add_chart(slide):
    data = CategoryChartData()
    data.categories = ["a", "b"]
    data.add_series("s", (1, 2))
    return slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1),
                                  Inches(1), Inches(4), Inches(3), data)


class TestMoveShapeToSlide(unittest.TestCase):
    def test_text_box_moves_with_its_geometry(self):
        pres = blank_deck()
        source, target = pres.slides
        box = source.shapes.add_textbox(Inches(2), Inches(3), Inches(4), Inches(1))
        box.text_frame.text = "moved"
        index = ppt_utils.move_shape_to_slide(source, box, target)
        self.assertEqual(len(source.shapes), 0)
        moved = target.shapes[index]
        self.assertEqual(moved.text_frame.text, "moved")
        self.assertEqual((moved.left, moved.top), (Inches(2), Inches(3)))

    def test_picture_keeps_resolving_after_save(self):
        pres = blank_deck()
        source, target = pres.slides
        picture = source.shapes.add_picture(io.BytesIO(PNG), Inches(1), Inches(1))
        ppt_utils.move_shape_to_slide(source, picture, target)
        # The source no longer references the image; the target does.
        self.assertFalse(any(r.reltype.endswith("/image")
                             for r in source.part.rels.values()))
        again = reopen(pres)
        self.assertEqual(len(again.slides[0].shapes), 0)
        self.assertEqual(again.slides[1].shapes[0].image.blob, PNG)

    def test_chart_moves_and_stays_editable(self):
        pres = blank_deck()
        source, target = pres.slides
        frame = add_chart(source)
        ppt_utils.move_shape_to_slide(source, frame, target)
        again = reopen(pres)
        moved = again.slides[1].shapes[0]
        self.assertTrue(moved.has_chart)
        self.assertEqual(list(moved.chart.plots[0].categories), ["a", "b"])
        # No dangling rId, no orphaned part left behind on the source.
        report = deck_validation.validate_presentation(pres)
        self.assertEqual(report["counts"], {"error": 0, "warning": 0, "info": 0})

    def test_shape_ids_are_unique_on_the_target(self):
        pres = blank_deck()
        source, target = pres.slides
        target.shapes.add_textbox(0, 0, Inches(1), Inches(1))
        box = source.shapes.add_textbox(0, 0, Inches(1), Inches(1))
        ppt_utils.move_shape_to_slide(source, box, target)
        ids = [s.shape_id for s in target.shapes]
        self.assertEqual(len(ids), len(set(ids)))

    def test_placeholder_text_moves_into_the_matching_placeholder(self):
        pres = Presentation()
        source = pres.slides.add_slide(pres.slide_layouts[1])
        target = pres.slides.add_slide(pres.slide_layouts[1])
        body = source.placeholders[1]
        body.text_frame.text = "belongs on the next slide"
        index = ppt_utils.move_shape_to_slide(source, body, target)
        self.assertEqual(target.shapes[index].text_frame.text,
                         "belongs on the next slide")
        self.assertEqual(body.text_frame.text, "")

    def test_placeholder_without_a_match_is_refused(self):
        pres = Presentation()
        source = pres.slides.add_slide(pres.slide_layouts[1])
        target = pres.slides.add_slide(pres.slide_layouts[6])  # blank
        source.placeholders[1].text_frame.text = "x"
        with self.assertRaises(ValueError):
            ppt_utils.move_shape_to_slide(source, source.placeholders[1], target)


class TestDeleteAndRestack(unittest.TestCase):
    def test_delete_drops_an_unshared_image_rel_only(self):
        pres = blank_deck(1)
        slide = pres.slides[0]
        first = slide.shapes.add_picture(io.BytesIO(PNG), 0, 0)
        second = slide.shapes.add_picture(io.BytesIO(PNG), Inches(2), 0)
        # python-pptx reuses one rId for the same image.
        self.assertEqual(first._element.blip_rId, second._element.blip_rId)
        ppt_utils.delete_shape(slide, first)
        self.assertIn(second._element.blip_rId, slide.part.rels)
        ppt_utils.delete_shape(slide, second)
        self.assertFalse(any(r.reltype.endswith("/image")
                             for r in slide.part.rels.values()))

    def test_z_order(self):
        pres = blank_deck(1)
        slide = pres.slides[0]
        a = slide.shapes.add_textbox(0, 0, 10, 10)
        b = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, 10, 10)
        ppt_utils.set_shape_z_order(slide, a, "front")
        self.assertEqual(list(slide.shapes)[-1].shape_id, a.shape_id)
        ppt_utils.set_shape_z_order(slide, a, "back")
        self.assertEqual(list(slide.shapes)[0].shape_id, a.shape_id)
        self.assertEqual(list(slide.shapes)[1].shape_id, b.shape_id)

    def test_reorder_slides(self):
        pres = blank_deck(3)
        for i, slide in enumerate(pres.slides):
            slide.shapes.add_textbox(0, 0, 10, 10).text_frame.text = str(i)
        ppt_utils.reorder_slides(pres, [2, 0, 1])
        texts = [s.shapes[0].text_frame.text for s in reopen(pres).slides]
        self.assertEqual(texts, ["2", "0", "1"])
        with self.assertRaises(ValueError):
            ppt_utils.reorder_slides(pres, [0, 0, 1])


class _FakeApp:
    def __init__(self):
        self.tools = {}

    def tool(self, **kwargs):
        def decorator(fn):
            self.tools[fn.__name__] = fn
            return fn
        return decorator


class TestManageShapeTool(unittest.TestCase):
    def setUp(self):
        self.store = PresentationStore(ttl_seconds=60, max_items=10)
        self.pid = self.store.new_id()
        pres = Presentation()
        pres.slides.add_slide(pres.slide_layouts[1])
        pres.slides.add_slide(pres.slide_layouts[6])
        pres.slides[0].shapes.add_textbox(
            Inches(1), Inches(1), Inches(2), Inches(1)).text_frame.text = "card"
        self.store[self.pid] = pres
        app = _FakeApp()
        register_slide_tools(app, self.store)
        self.tool = app.tools["manage_shape"]
        self.pres = pres

    def test_delete(self):
        before = len(self.pres.slides[0].shapes)
        result = self.tool(self.pid, 0, before - 1, "delete")
        self.assertNotIn("error", result)
        self.assertEqual(len(self.pres.slides[0].shapes), before - 1)

    def test_deleting_a_placeholder_empties_it(self):
        self.pres.slides[0].placeholders[1].text_frame.text = "prompt"
        result = self.tool(self.pid, 0, 1, "delete")
        self.assertIn("Cleared placeholder", result["message"])
        self.assertEqual(self.pres.slides[0].placeholders[1].text_frame.text, "")

    def test_move_to_slide(self):
        index = len(self.pres.slides[0].shapes) - 1
        result = self.tool(self.pid, 0, index, "move_to_slide",
                           target_slide_index=1, left=3.0)
        self.assertEqual(result["slide_index"], 1)
        moved = self.pres.slides[1].shapes[result["shape_index"]]
        self.assertEqual(moved.text_frame.text, "card")
        self.assertEqual(moved.left, Inches(3))

    def test_set_geometry(self):
        index = len(self.pres.slides[0].shapes) - 1
        self.tool(self.pid, 0, index, "set_geometry", width=5.0)
        self.assertEqual(self.pres.slides[0].shapes[index].width, Inches(5))
        self.assertIn("error", self.tool(self.pid, 0, index, "set_geometry"))

    def test_errors_are_actionable(self):
        self.assertIn("get_slide_info",
                      self.tool(self.pid, 0, 99, "delete")["error"])
        self.assertIn("Unknown operation",
                      self.tool(self.pid, 0, 0, "explode")["error"])
        self.assertIn("error", self.tool(self.pid, 0, 0, "move_to_slide"))

    def test_slide_info_shows_text_and_inches(self):
        info = ppt_utils.get_slide_info(self.pres.slides[0], 0)
        card = info["shapes"][-1]
        self.assertEqual(card["text"], "card")
        self.assertEqual(card["left_in"], 1.0)


if __name__ == "__main__":
    unittest.main()
