"""Tests for the server-side asset library: list_assets and
add_asset_to_slide, including the name checks that keep the tool from
reading files outside the library folder."""
import os
import sys
import tempfile
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from pptx import Presentation

from state import PresentationStore
from tools.image_tools import register_image_tools
from tests.test_image_tools import _FakeApp, png_bytes, EMU_PER_INCH


class AssetToolTestCase(unittest.TestCase):
    def setUp(self):
        self._saved = {k: os.environ.get(k)
                       for k in ("PPT_ASSET_PATH", "DIAL_IMAGE_MAX_MB")}
        os.environ.pop("DIAL_IMAGE_MAX_MB", None)
        self.tmp = tempfile.TemporaryDirectory()
        self.library = os.path.join(self.tmp.name, "icons")
        os.mkdir(self.library)
        for name in ("brand_email_blue.png", "brand_email_outline.png",
                     "brand_folder_blue.png"):
            with open(os.path.join(self.library, name), "wb") as fh:
                fh.write(png_bytes(800, 800))
        # Not listable: hidden, not an image, a directory.
        for name in (".DS_Store", "manifest.json"):
            with open(os.path.join(self.library, name), "w") as fh:
                fh.write("{}")
        os.mkdir(os.path.join(self.library, "nested.png"))
        # Outside the library, next to it.
        with open(os.path.join(self.tmp.name, "secret.png"), "wb") as fh:
            fh.write(png_bytes(10, 10))
        os.environ["PPT_ASSET_PATH"] = self.library

        self.store = PresentationStore(ttl_seconds=60, max_items=10)
        self.pid = self.store.new_id()
        pres = Presentation()
        pres.slides.add_slide(pres.slide_layouts[6])
        self.store[self.pid] = pres
        app = _FakeApp()
        register_image_tools(app, self.store)
        self.list_assets = app.tools["list_assets"]
        self.add_asset = app.tools["add_asset_to_slide"]

    def tearDown(self):
        self.tmp.cleanup()
        for key, value in self._saved.items():
            if value is None:
                os.environ.pop(key, None)
            else:
                os.environ[key] = value

    def add(self, **kwargs):
        kwargs.setdefault("presentation_id", self.pid)
        kwargs.setdefault("slide_index", 0)
        kwargs.setdefault("name", "brand_email_blue.png")
        return self.add_asset(**kwargs)


class TestListAssets(AssetToolTestCase):
    def test_lists_only_images_sorted(self):
        result = self.list_assets()
        self.assertEqual(result["assets"], ["brand_email_blue.png",
                                            "brand_email_outline.png",
                                            "brand_folder_blue.png"])
        self.assertEqual(result["count"], 3)

    def test_query_filters_case_insensitively(self):
        result = self.list_assets(query="EMAIL")
        self.assertEqual(result["assets"], ["brand_email_blue.png",
                                            "brand_email_outline.png"])

    def test_no_match_says_what_to_do(self):
        result = self.list_assets(query="rocket")
        self.assertEqual(result["assets"], [])
        self.assertIn("render_svg_icon", result["note"])

    def test_unconfigured_library_is_an_error_not_a_crash(self):
        os.environ.pop("PPT_ASSET_PATH")
        self.assertIn("PPT_ASSET_PATH", self.list_assets()["error"])

    def test_missing_folder_is_an_error_not_a_crash(self):
        os.environ["PPT_ASSET_PATH"] = os.path.join(self.tmp.name, "gone")
        self.assertIn("not available", self.list_assets()["error"])


class TestAddAsset(AssetToolTestCase):
    def test_places_icon_in_a_square_box(self):
        result = self.add(left=1, top=2, width=0.8, height=0.8)
        self.assertNotIn("error", result)
        self.assertEqual(result["asset"], "brand_email_blue.png")
        self.assertEqual(result["placed"], {"left": 1.0, "top": 2.0,
                                            "width": 0.8, "height": 0.8})
        pic = self.store[self.pid].slides[0].shapes[result["shape_index"]]
        self.assertEqual(pic.width, int(0.8 * EMU_PER_INCH))

    def test_contain_keeps_the_aspect_ratio(self):
        result = self.add(left=0, top=0, width=2, height=1)
        self.assertEqual(result["placed"]["width"], 1.0)
        self.assertEqual(result["placed"]["left"], 0.5)

    def test_unknown_name_points_to_list_assets(self):
        result = self.add(name="brand_rocket_blue.png")
        self.assertIn("list_assets", result["error"])

    def test_paths_are_refused(self):
        for name in ("../secret.png", os.path.join(self.tmp.name, "secret.png"),
                     "nested.png/x.png", ".DS_Store", "manifest.json",
                     "nested.png", ""):
            with self.subTest(name=name):
                self.assertIn("error", self.add(name=name))
        self.assertEqual(len(self.store[self.pid].slides[0].shapes), 0)

    def test_shared_argument_checks_apply(self):
        self.assertIn("Unknown", self.add(presentation_id="nope")["error"])
        self.assertIn("slide index", self.add(slide_index=5)["error"])
        self.assertIn("fit", self.add(fit="squash")["error"])

    def test_oversize_asset_refused(self):
        os.environ["DIAL_IMAGE_MAX_MB"] = "0.000001"
        self.assertIn("different asset", self.add()["error"])


if __name__ == "__main__":
    unittest.main()
