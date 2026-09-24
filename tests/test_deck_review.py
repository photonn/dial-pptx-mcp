"""Tests for the deck-level review: the slide outline and its use by the
visual QA loop (coherence issues, batching, the pass rule, the result)."""
import os
import sys
import threading
import unittest
from unittest.mock import patch
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt

import deck_review
import visual_fix
import visual_qa

ENV = {
    "VISION_LLM_ENDPOINT": "https://example.invalid/responses",
    "VISION_LLM_API_KEY": "k",
    "VISION_LLM_MODEL": "m",
}
MANAGED = list(ENV) + ["VISUAL_QA_COHERENCE", "VISION_LLM_BATCH_SLIDES",
                       "VISION_LLM_MAX_SLIDES", "VISUAL_QA_MAX_ITERATIONS"]


def titled_deck(titles):
    """One "Title and Content" slide per title, body left empty."""
    pres = Presentation()
    for title in titles:
        slide = pres.slides.add_slide(pres.slide_layouts[1])
        slide.shapes.title.text_frame.text = title
    return pres


class EnvCase(unittest.TestCase):
    def setUp(self):
        self._saved = {k: os.environ.get(k) for k in MANAGED}
        for k in MANAGED:
            os.environ.pop(k, None)
        os.environ.update(ENV)

    def tearDown(self):
        for k, v in self._saved.items():
            if v is None:
                os.environ.pop(k, None)
            else:
                os.environ[k] = v


class TestOutline(unittest.TestCase):
    def test_title_body_and_empty_slide(self):
        pres = titled_deck(["Agenda", "Revenue"])
        pres.slides[0].placeholders[1].text_frame.text = "1 Revenue"
        outline = deck_review.slide_outline(pres)
        self.assertEqual(outline[0]["title"], "Agenda")
        self.assertEqual(outline[0]["body_elements"], 1)
        # The second slide has a title and nothing else: the blank-slide
        # signal the coherence reviewer is told to treat as critical.
        self.assertEqual(outline[1]["body_elements"], 0)
        body = [e for e in outline[1]["elements"]
                if e["kind"] == "placeholder:object"
                or e["kind"] == "placeholder:body"]
        self.assertTrue(body and body[0].get("empty"))

    def test_text_covered_by_a_later_card_is_flagged(self):
        pres = titled_deck(["Section"])
        slide = pres.slides[0]
        title = slide.shapes.title
        slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, title.left, title.top,
                               title.width, title.height)
        outline = deck_review.slide_outline(pres)
        entry = next(e for e in outline[0]["elements"] if e.get("role") == "title")
        self.assertEqual(entry["drawn_over_by"], [len(slide.shapes) - 1])

    def test_unfilled_text_box_does_not_hide_anything(self):
        pres = titled_deck(["Section"])
        slide = pres.slides[0]
        title = slide.shapes.title
        slide.shapes.add_textbox(title.left, title.top, title.width, title.height)
        outline = deck_review.slide_outline(pres)
        entry = next(e for e in outline[0]["elements"] if e.get("role") == "title")
        self.assertNotIn("drawn_over_by", entry)

    def test_source_line_is_not_body_content(self):
        pres = titled_deck(["Chart"])
        slide = pres.slides[0]
        box = slide.shapes.add_textbox(Inches(0.5), pres.slide_height - Inches(0.4),
                                       Inches(6), Inches(0.3))
        box.text_frame.text = "Source: annual report"
        outline = deck_review.slide_outline(pres)
        self.assertEqual(outline[0]["body_elements"], 0)
        self.assertEqual(outline[0]["elements"][-1]["role"], "footnote")

    def test_prompt_carries_the_outline_and_the_agenda_check(self):
        prompt = deck_review.coherence_prompt(
            deck_review.slide_outline(titled_deck(["Agenda", "Revenue"])))
        self.assertIn('"title":"Revenue"', prompt)
        for phrase in ("Agenda / index", "belongs to ANOTHER slide",
                       "body_elements 0", "Never invent"):
            self.assertIn(phrase, prompt)

    def test_switch(self):
        with patch.dict(os.environ, {"VISUAL_QA_COHERENCE": "false"}):
            self.assertFalse(deck_review.coherence_enabled())
        with patch.dict(os.environ, {}, clear=False):
            os.environ.pop("VISUAL_QA_COHERENCE", None)
            self.assertTrue(deck_review.coherence_enabled())


class TestReviewPrompt(unittest.TestCase):
    def test_checklist_covers_the_failures_seen_in_the_field(self):
        prompt = visual_qa.review_prompt(False)
        for phrase in ("only its title", "does not match the slide title",
                       "Text showing through", "Missing-glyph", "Doubled bullets",
                       "squashed", "Distorted aspect ratio", "header rule",
                       '"category"'):
            self.assertIn(phrase, prompt)

    def test_inventory_is_attached_per_slide(self):
        pres = titled_deck(["Agenda", "Revenue"])
        inventory = deck_review.slide_outline(pres, [2])
        prompt = visual_qa.review_prompt(False, slides=[2], inventory=inventory)
        self.assertIn("image 1 = slide 2", prompt)
        self.assertIn("slide 2 (layout", prompt)
        self.assertIn("'Revenue'", prompt)

    def test_focus_is_added_not_substituted(self):
        prompt = visual_qa.review_prompt(False, focus="check the logo")
        self.assertIn("never instead of", prompt)
        self.assertIn("Empty, incomplete or wrong content", prompt)

    def test_images_are_sent_at_high_detail(self):
        with patch.dict(os.environ, ENV):
            payload = visual_qa.VisionLLM().build_payload([b"png"], "p")
        self.assertEqual(payload["input"][0]["content"][1]["detail"], "high")


class TestBatchedReview(EnvCase):
    def _deck(self, n):
        return titled_deck([f"Slide {i}" for i in range(1, n + 1)])

    def test_whole_deck_is_split_into_batches(self):
        os.environ["VISION_LLM_BATCH_SLIDES"] = "3"
        os.environ["VISUAL_QA_COHERENCE"] = "false"
        prompts, lock = [], threading.Lock()

        def fake_review(self_llm, images, prompt, timeout=None):
            with lock:
                prompts.append((len(images), prompt))
            return {"passed": True, "issues": []}

        with patch.object(visual_qa, "_render_deck",
                          return_value=[b"png"] * 8), \
             patch.object(visual_qa.VisionLLM, "review", fake_review):
            verdict = visual_qa.inspect_presentation(self._deck(8))
        self.assertEqual(sorted(n for n, _ in prompts), [2, 3, 3])
        self.assertTrue(any("image 1 = slide 7" in p for _, p in prompts))
        self.assertEqual(verdict["slides_reviewed"], 8)
        self.assertTrue(verdict["passed"])

    def test_model_pass_flag_does_not_override_a_major_issue(self):
        os.environ["VISUAL_QA_COHERENCE"] = "false"
        verdict = {"passed": True, "issues": [
            {"slide": 1, "severity": "major", "description": "clipped"}]}
        with patch.object(visual_qa, "_render_deck", return_value=[b"png"]), \
             patch.object(visual_qa.VisionLLM, "review",
                          lambda *a, **k: dict(verdict, issues=list(verdict["issues"]))):
            result = visual_qa.inspect_presentation(self._deck(1))
        self.assertFalse(result["passed"])

    def test_issue_for_a_slide_outside_its_batch_is_dropped(self):
        os.environ["VISUAL_QA_COHERENCE"] = "false"
        verdict = {"passed": False, "issues": [
            {"slide": 9, "severity": "major", "description": "x"}]}
        with patch.object(visual_qa, "_render_deck", return_value=[b"png"] * 2), \
             patch.object(visual_qa.VisionLLM, "review",
                          lambda *a, **k: dict(verdict, issues=list(verdict["issues"]))):
            result = visual_qa.inspect_presentation(self._deck(2))
        self.assertEqual(result["issues"], [])

    def test_slides_beyond_the_cap_are_reported_not_passed(self):
        os.environ["VISION_LLM_MAX_SLIDES"] = "2"
        os.environ["VISUAL_QA_COHERENCE"] = "false"
        seen = {}

        def fake_render(pres, max_slides=None, slides=None):
            seen["slides"] = slides
            return [b"png"] * (len(slides) if slides else 3)

        with patch.object(visual_qa, "_render_deck", fake_render), \
             patch.object(visual_qa.VisionLLM, "review",
                          lambda *a, **k: {"passed": True, "issues": []}):
            outcome = visual_qa.inspect_and_repair(self._deck(3))
        self.assertEqual(seen["slides"], [1, 2])
        self.assertEqual(outcome["slides_not_reviewed"], [3])
        self.assertIn("VISION_LLM_MAX_SLIDES", outcome["review_note"])


class TestCallerContext(EnvCase):
    def test_review_threads_see_the_callers_request_context(self):
        """The DIAL provider reads the caller's credentials from the MCP
        SDK's request contextvar; the parallel review threads must see it or
        every DIAL-routed review fails with "No DIAL credentials"."""
        from mcp.server.lowlevel.server import request_ctx

        os.environ["VISION_LLM_BATCH_SLIDES"] = "1"
        seen, lock = [], threading.Lock()

        def fake_review(self_llm, images, prompt, timeout=None):
            with lock:
                seen.append(request_ctx.get(None))
            return {"passed": True, "issues": []}

        marker = object()
        token = request_ctx.set(marker)
        try:
            with patch.object(visual_qa, "_render_deck",
                              return_value=[b"png"] * 3), \
                 patch.object(visual_qa.VisionLLM, "review", fake_review):
                visual_qa.inspect_presentation(titled_deck(["A", "B", "C"]))
        finally:
            request_ctx.reset(token)
        # Three visual batches plus the coherence call, all with the context.
        self.assertEqual(len(seen), 4)
        self.assertTrue(all(value is marker for value in seen))


class TestCoherenceInTheLoop(EnvCase):
    def _script(self, visual, coherence):
        """Fake LLM: image-less calls are the coherence review."""
        calls = {"visual": 0, "coherence": 0}

        def fake_review(self_llm, images, prompt, timeout=None):
            key = "visual" if images else "coherence"
            script = visual if images else coherence
            verdict = script[min(calls[key], len(script) - 1)]
            calls[key] += 1
            if isinstance(verdict, Exception):
                raise verdict
            return {"passed": verdict["passed"],
                    "issues": [dict(i) for i in verdict["issues"]]}
        return fake_review, calls

    def test_misplaced_content_is_moved_and_the_deck_passes(self):
        pres = titled_deck(["Agenda", "Section", "Scale, focus, launch"])
        card = pres.slides[1].shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(1), Inches(2), Inches(3), Inches(2))
        card.text_frame.text = "SCALE"
        card_index = len(pres.slides[1].shapes) - 1
        misplaced = {"slide": 2, "severity": "critical",
                     "category": "misplaced_content", "related_slides": [3],
                     "description": "cards belong on slide 3"}
        fake_review, calls = self._script(
            visual=[{"passed": True, "issues": []}],
            coherence=[{"passed": False, "issues": [misplaced]},
                       {"passed": True, "issues": []}])
        planned = {}

        def fake_plan(llm, issues, pres_, images, image_slides=None,
                      author_actions=None):
            planned["issues"] = issues
            return [{"op": "move_shape_to_slide", "slide": 2,
                     "shape_index": card_index, "target_slide": 3}]

        with patch.object(visual_qa, "_render_deck",
                          side_effect=lambda p, m=None, slides=None:
                          [b"png"] * (len(slides) if slides else 3)), \
             patch.object(visual_qa.VisionLLM, "review", fake_review), \
             patch.object(visual_fix, "plan_repairs", fake_plan):
            outcome = visual_qa.inspect_and_repair(pres)
        self.assertTrue(outcome["passed"])
        self.assertEqual(planned["issues"][0]["check"], "coherence")
        self.assertEqual(outcome["checks"], ["visual", "coherence"])
        self.assertEqual(outcome["repair_rounds"][0]["changes"],
                         [f"slide 2 #{card_index} move_shape_to_slide -> slide 3"])
        self.assertEqual(pres.slides[2].shapes[-1].text_frame.text, "SCALE")
        self.assertEqual(calls["coherence"], 2)  # re-checked after the move

    def test_missing_content_comes_back_as_an_author_action(self):
        pres = titled_deck(["Agenda", "Revenue", "Costs"])
        empty = {"slide": 3, "severity": "critical", "category": "empty_slide",
                 "description": "blank", "suggested_fix": "add the cost chart"}
        fake_review, _ = self._script(
            visual=[{"passed": True, "issues": []}],
            coherence=[{"passed": False, "issues": [empty]}])

        def fake_plan(llm, issues, pres_, images, image_slides=None,
                      author_actions=None):
            author_actions.append({"slide": 3, "action": "build the cost chart"})
            return []

        with patch.object(visual_qa, "_render_deck", return_value=[b"png"] * 3), \
             patch.object(visual_qa.VisionLLM, "review", fake_review), \
             patch.object(visual_fix, "plan_repairs", fake_plan):
            outcome = visual_qa.inspect_and_repair(pres)
        self.assertFalse(outcome["passed"])
        self.assertEqual(outcome["action_required"],
                         [{"slide": 3, "action": "build the cost chart"}])
        self.assertEqual(outcome["issues"][0]["category"], "empty_slide")

    def test_a_failed_coherence_call_degrades_to_the_visual_review(self):
        pres = titled_deck(["A", "B", "C"])
        fake_review, _ = self._script(
            visual=[{"passed": True, "issues": []}],
            coherence=[visual_qa.VisualQAError("HTTP 500")])
        with patch.object(visual_qa, "_render_deck", return_value=[b"png"] * 3), \
             patch.object(visual_qa.VisionLLM, "review", fake_review):
            outcome = visual_qa.inspect_and_repair(pres)
        self.assertTrue(outcome["passed"])

    def test_scoped_calls_skip_the_story_review(self):
        pres = titled_deck(["A", "B", "C"])
        fake_review, calls = self._script(
            visual=[{"passed": True, "issues": []}],
            coherence=[{"passed": False, "issues": []}])
        with patch.object(visual_qa, "_render_deck", return_value=[b"png"]), \
             patch.object(visual_qa.VisionLLM, "review", fake_review):
            outcome = visual_qa.inspect_and_repair(pres, slides=[2])
        self.assertEqual(calls["coherence"], 0)
        self.assertEqual(outcome["checks"], ["visual"])

    def test_minor_findings_are_reported_on_a_pass(self):
        pres = titled_deck(["A"])
        minor = {"slide": 1, "severity": "minor", "description": "nit"}
        with patch.object(visual_qa, "_render_deck", return_value=[b"png"]), \
             patch.object(visual_qa.VisionLLM, "review",
                          lambda *a, **k: {"passed": True,
                                           "issues": [dict(minor)]}):
            outcome = visual_qa.inspect_and_repair(pres)
        self.assertTrue(outcome["passed"])
        self.assertEqual(outcome["minor_issues"][0]["description"], "nit")


class TestNewRepairOperations(unittest.TestCase):
    def _deck(self):
        pres = Presentation()
        for _ in range(3):
            slide = pres.slides.add_slide(pres.slide_layouts[6])
            for i in range(3):
                slide.shapes.add_textbox(Inches(i), 0, Inches(1), Inches(1)
                                         ).text_frame.text = f"t{i}"
        return pres

    def test_indexes_refer_to_the_slide_as_described(self):
        pres = self._deck()
        result = visual_fix.apply_repairs(pres, [
            {"op": "delete_shape", "slide": 1, "shape_index": 0},
            {"op": "set_text", "slide": 1, "shape_index": 2, "text": "last"},
            {"op": "set_text", "slide": 1, "shape_index": 0, "text": "gone"}])
        texts = [s.text_frame.text for s in pres.slides[0].shapes]
        # Index 2 still meant the original third box, not the (now) third
        # position; the deleted shape cannot be addressed again.
        self.assertEqual(texts, ["t1", "last"])
        self.assertIn("already deleted", result["skipped"][0]["reason"])

    def test_move_to_slide_respects_scope(self):
        pres = self._deck()
        result = visual_fix.apply_repairs(pres, [
            {"op": "move_shape_to_slide", "slide": 1, "shape_index": 0,
             "target_slide": 3}], allowed_slides=[1])
        self.assertEqual(result["skipped"][0]["reason"], "slide out of scope")

    def test_reorder_runs_last_and_is_flagged(self):
        pres = self._deck()
        pres.slides[2].shapes[0].text_frame.text = "was third"
        result = visual_fix.apply_repairs(pres, [
            {"op": "reorder_slides", "order": [3, 1, 2]},
            # Addresses the pre-reorder numbering even though listed after.
            {"op": "set_text", "slide": 3, "shape_index": 1, "text": "edited"}])
        self.assertTrue(result["slides_reordered"])
        first = pres.slides[0].shapes
        self.assertEqual((first[0].text_frame.text, first[1].text_frame.text),
                         ("was third", "edited"))

    def test_reorder_is_validated(self):
        pres = self._deck()
        for bad in ([1, 2], [1, 1, 2], "3,1,2"):
            result = visual_fix.apply_repairs(
                pres, [{"op": "reorder_slides", "order": bad}])
            self.assertEqual(result["skipped"][0]["reason"], "bad slide order")
        result = visual_fix.apply_repairs(
            pres, [{"op": "reorder_slides", "order": [2, 1, 3]}],
            allowed_slides=[1, 2])
        self.assertEqual(result["skipped"][0]["reason"], "slide out of scope")

    def test_color_clear_and_restack(self):
        pres = self._deck()
        result = visual_fix.apply_repairs(pres, [
            {"op": "set_font_color", "slide": 1, "shape_index": 0,
             "color": "#0055AA"},
            {"op": "set_font_color", "slide": 1, "shape_index": 1,
             "color": "blue"},
            {"op": "clear_text", "slide": 1, "shape_index": 2},
            {"op": "send_to_back", "slide": 2, "shape_index": 2}])
        shapes = pres.slides[0].shapes
        self.assertEqual(str(shapes[0].text_frame.paragraphs[0].runs[0]
                             .font.color.rgb), "0055AA")
        self.assertEqual(shapes[2].text_frame.text, "")
        self.assertEqual(pres.slides[1].shapes[0].text_frame.text, "t2")
        self.assertEqual([s["reason"] for s in result["skipped"]], ["bad color"])

    def test_planner_sees_related_slides_and_returns_author_actions(self):
        pres = self._deck()
        captured = {}

        class FakeLLM:
            def ask_json(self, images, prompt):
                captured.update(images=images, prompt=prompt)
                return {"operations": [],
                        "author_actions": [{"slide": 2, "action": "add data"},
                                           "junk"]}

        actions = []
        visual_fix.plan_repairs(
            FakeLLM(), [{"slide": 1, "related_slides": [3],
                         "description": "cards belong on slide 3"}],
            pres, [b"i1", b"i2", b"i3"], [1, 2, 3], author_actions=actions)
        self.assertEqual(captured["images"], [b"i1", b"i3"])
        self.assertIn('"slide": 3', captured["prompt"])
        self.assertEqual(actions, [{"slide": 2, "action": "add data"}])

    def test_repair_prompt_advertises_the_structural_operations(self):
        for op in ("move_shape_to_slide", "reorder_slides", "bring_to_front",
                   "send_to_back", "clear_text", "set_font_color",
                   "author_actions", "Never invent"):
            self.assertIn(op, visual_fix.REPAIR_PROMPT)


if __name__ == "__main__":
    unittest.main()
