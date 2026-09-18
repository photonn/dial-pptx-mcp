"""Tests for what the repair loop does *before* it asks the model anything:
seeding it with a verdict the caller already has, sizing its budget by scope,
and folding in the structural findings deck_validation can produce without a
render.
"""
import os
import sys
import unittest
from unittest.mock import patch
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from pptx import Presentation
from pptx.util import Inches

import visual_fix
import visual_qa
from state import PresentationStore
from tools.visual_tools import register_visual_tools

VISION_ENV = {
    "VISION_LLM_ENDPOINT": "https://example.invalid/responses",
    "VISION_LLM_API_KEY": "k",
    "VISION_LLM_MODEL": "m",
}
BUDGET_VARS = ("VISUAL_QA_MAX_ITERATIONS", "VISUAL_QA_MAX_ITERATIONS_SLIDE",
               "VISUAL_QA_REPAIR_SEVERITY", "VISUAL_QA_ENFORCE",
               "VISION_LLM_MODEL_REVIEW", "VISION_LLM_MODEL_PLAN")


def make_deck(slides=3):
    pres = Presentation()
    for i in range(slides):
        slide = pres.slides.add_slide(pres.slide_layouts[6])
        box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4),
                                       Inches(1))
        box.text_frame.text = f"Slide {i + 1}"
    return pres


class _FakeApp:
    def __init__(self):
        self.tools = {}

    def tool(self, **kwargs):
        def decorator(fn):
            self.tools[fn.__name__] = fn
            return fn
        return decorator


class EnvTestCase(unittest.TestCase):
    def setUp(self):
        keys = list(VISION_ENV) + list(BUDGET_VARS)
        self._saved = {k: os.environ.get(k) for k in keys}
        for k in keys:
            os.environ.pop(k, None)
        os.environ.update(VISION_ENV)

    def tearDown(self):
        for k, v in self._saved.items():
            if v is None:
                os.environ.pop(k, None)
            else:
                os.environ[k] = v


class _Loop(EnvTestCase):
    """Drives inspect_and_repair with a scripted reviewer and planner."""

    def run_loop(self, verdicts, plans, **kwargs):
        calls = {"review": 0, "plan": 0, "render": 0}

        def fake_render(pres, max_slides=None, slides=None, dpi=None):
            calls["render"] += 1
            return [b"\x89PNG-fake"] * (len(slides) if slides else 3)

        def fake_review(self_llm, images, prompt, timeout=None):
            v = verdicts[min(calls["review"], len(verdicts) - 1)]
            calls["review"] += 1
            return dict(v)

        def fake_plan(llm, issues, pres, images, image_slides=None):
            p = plans[min(calls["plan"], len(plans) - 1)]
            calls["plan"] += 1
            return p

        with patch.object(visual_qa, "_render_deck", fake_render), \
             patch.object(visual_qa.VisionLLM, "review", fake_review), \
             patch.object(visual_fix, "plan_repairs", fake_plan):
            outcome = visual_qa.inspect_and_repair(make_deck(), **kwargs)
        return outcome, calls


class TestSeededRepair(_Loop):
    ONE_ISSUE = [{"slide": 2, "severity": "major", "description": "overflow",
                  "suggested_fix": "shrink it"}]
    PLAN = [{"op": "set_font_size", "slide": 2, "shape_index": 0,
             "size_pt": 14}]

    def test_a_supplied_verdict_replaces_the_first_review(self):
        outcome, calls = self.run_loop(
            [{"passed": True, "issues": []}], [self.PLAN], slides=[2],
            initial_verdict={"passed": False, "issues": self.ONE_ISSUE})
        self.assertTrue(outcome["passed"])
        # One review, not two: the caller already paid for the first one.
        self.assertEqual(calls["review"], 1)
        self.assertEqual(calls["plan"], 1)

    def test_without_the_seed_the_same_work_costs_an_extra_review(self):
        outcome, calls = self.run_loop(
            [{"passed": False, "issues": self.ONE_ISSUE},
             {"passed": True, "issues": []}], [self.PLAN], slides=[2])
        self.assertTrue(outcome["passed"])
        self.assertEqual(calls["review"], 2)

    def test_a_seed_that_needs_no_repair_renders_nothing(self):
        outcome, calls = self.run_loop(
            [{"passed": True, "issues": []}], [[]], slides=[2],
            initial_verdict={"passed": True, "issues": []})
        self.assertTrue(outcome["passed"])
        self.assertEqual(calls["review"], 0)
        self.assertEqual(calls["render"], 0)

    def test_a_seed_cannot_reach_outside_the_scope(self):
        seeded = []

        def capture(llm, issues, pres, images, image_slides=None):
            seeded.extend(issues)
            return []

        with patch.object(visual_qa, "_render_deck",
                          return_value=[b"\x89PNG-fake"]), \
             patch.object(visual_qa.VisionLLM, "review",
                          lambda *a, **k: {"passed": True, "issues": []}), \
             patch.object(visual_fix, "plan_repairs", capture):
            visual_qa.inspect_and_repair(
                make_deck(), slides=[2],
                initial_verdict={"passed": False, "issues": [
                    {"slide": 1, "severity": "major", "description": "elsewhere"},
                    {"slide": 2, "severity": "major", "description": "in scope"}]})
        self.assertEqual([i["description"] for i in seeded], ["in scope"])

    def test_a_malformed_seed_falls_back_to_a_normal_first_review(self):
        for bad in ("not a dict", {"issues": []}, {"passed": False,
                                                   "issues": "nope"}):
            with self.subTest(seed=bad):
                outcome, calls = self.run_loop(
                    [{"passed": True, "issues": []}], [[]], slides=[2],
                    initial_verdict=bad)
                self.assertTrue(outcome["passed"])
                self.assertEqual(calls["review"], 1)

    def test_the_tool_passes_its_issues_argument_through(self):
        store = PresentationStore(ttl_seconds=60, max_items=10)
        pid = store.new_id()
        store[pid] = make_deck()
        app = _FakeApp()
        register_visual_tools(app, store)
        seen = {}

        def fake_inspect(pres, slides=None, focus=None, max_iterations=None,
                         initial_verdict=None):
            seen["seed"] = initial_verdict
            return {"passed": True, "iterations": 1, "repair_rounds": []}

        with patch.object(visual_qa, "inspect_and_repair", fake_inspect):
            app.tools["visual_repair_slides"](pid, slides=[2],
                                              issues=self.ONE_ISSUE)
        self.assertEqual(seen["seed"], {"passed": False,
                                        "issues": self.ONE_ISSUE})


class TestIterationBudget(EnvTestCase):
    def test_scope_picks_the_default(self):
        # A deck-sized budget spent on one slide is how a 20s call becomes
        # a 200s one.
        self.assertEqual(visual_qa._iteration_budget(None), 10)
        self.assertEqual(visual_qa._iteration_budget([3]), 2)

    def test_explicit_argument_always_wins(self):
        os.environ["VISUAL_QA_MAX_ITERATIONS_SLIDE"] = "5"
        self.assertEqual(visual_qa._iteration_budget([3], 1), 1)
        self.assertEqual(visual_qa._iteration_budget([3]), 5)

    def test_environment_overrides(self):
        os.environ["VISUAL_QA_MAX_ITERATIONS"] = "4"
        self.assertEqual(visual_qa._iteration_budget(None), 4)
        os.environ["VISUAL_QA_MAX_ITERATIONS"] = "nonsense"
        self.assertEqual(visual_qa._iteration_budget(None), 10)
        self.assertEqual(visual_qa._iteration_budget(None, 0), 1)


class TestRoleModels(EnvTestCase):
    def test_roles_fall_back_to_the_single_model(self):
        self.assertEqual(visual_qa.VisionLLM("review").model, "m")
        self.assertEqual(visual_qa.VisionLLM("plan").model, "m")

    def test_review_can_be_routed_to_its_own_model(self):
        os.environ["VISION_LLM_MODEL_REVIEW"] = "fast-vision"
        self.assertEqual(visual_qa.VisionLLM("review").model, "fast-vision")
        self.assertEqual(visual_qa.VisionLLM("plan").model, "m")

    def test_reasoning_effort_is_only_sent_when_asked_for(self):
        llm = visual_qa.VisionLLM("review")
        self.assertNotIn("reasoning", llm.build_payload([], "p"))
        os.environ["VISION_LLM_REVIEW_REASONING"] = "low"
        try:
            self.assertEqual(
                visual_qa.VisionLLM("review").build_payload([], "p")["reasoning"],
                {"effort": "low"})
            # The plan role has its own variable and did not get this one.
            self.assertNotIn("reasoning",
                             visual_qa.VisionLLM("plan").build_payload([], "p"))
        finally:
            os.environ.pop("VISION_LLM_REVIEW_REASONING")


class TestStructuralFindings(EnvTestCase):
    """deck_validation numbers slides from 0, visual_qa from 1."""

    def off_slide_deck(self):
        pres = make_deck(2)
        shape = pres.slides[1].shapes[0]
        shape.left = -pres.slide_width
        return pres

    def test_slide_numbers_are_converted_to_one_based(self):
        issues, _ = visual_qa.structural_findings(self.off_slide_deck())
        self.assertEqual([i["slide"] for i in issues], [2])
        self.assertEqual(issues[0]["severity"], "critical")
        # The message deck_validation wrote counts from 0; what comes out of
        # here must count from 1, like everything else the planner reads.
        self.assertTrue(issues[0]["description"].startswith("Slide 2,"))
        self.assertIn("visual_repair_slides", issues[0]["suggested_fix"])

    def test_findings_outside_the_scope_are_dropped(self):
        pres = self.off_slide_deck()
        self.assertEqual(visual_qa.structural_findings(pres, [1])[0], [])
        self.assertEqual(len(visual_qa.structural_findings(pres, [2])[0]), 1)

    def test_a_distorted_picture_is_an_observation_not_an_issue(self):
        """Nothing in the repair whitelist can un-distort a picture, so it
        must not be able to buy a repair round."""
        problems = [{"code": "distorted_picture", "slide_index": 0,
                     "message": "Slide 0, shape 1 ('Pic') is stretched.",
                     "fix": "Re-add it.", "severity": "warning"}]
        with patch("deck_validation.validate_presentation",
                   return_value={"problems": problems}):
            issues, observations = visual_qa.structural_findings(make_deck(1))
        self.assertEqual(issues, [])
        self.assertEqual(observations[0]["slide"], 1)

    def test_the_loop_plans_from_validation_without_a_vision_finding(self):
        pres = self.off_slide_deck()
        planned = []

        def capture(llm, issues, pres_, images, image_slides=None):
            planned.append(list(issues))
            return []

        with patch.object(visual_qa, "_render_deck",
                          return_value=[b"\x89PNG-fake"]), \
             patch.object(visual_qa.VisionLLM, "review",
                          lambda *a, **k: {"passed": True, "issues": []}), \
             patch.object(visual_fix, "plan_repairs", capture):
            outcome = visual_qa.inspect_and_repair(pres, slides=[2])
        # The reviewer saw nothing wrong; the structural check did, so the
        # deck does not pass and the planner was given the finding.
        self.assertFalse(outcome["passed"])
        self.assertEqual(planned[0][0]["source"], "structure")

    def test_the_reviewer_is_told_what_the_structural_check_found(self):
        prompt = visual_qa.review_prompt(
            False, structural=[{"slide": 4, "description": "Slide 4, shape 0 "
                                                           "is off the slide"}])
        self.assertIn("do not report them again", prompt)
        self.assertIn("slide 4", prompt)


if __name__ == "__main__":
    unittest.main()
