"""Tests for the build-quality fixes: measured empty space, formatting-safe
text rewrites, flat styled shapes, zero-based charts with working data labels,
straight lines in validation, and the vision request's retry."""
import io
import os
import sys
import unittest
from pathlib import Path
from unittest.mock import patch

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import httpx
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

import deck_validation
import layout_space
import utils as ppt_utils
import visual_fix
import visual_qa

SLIDE_W_IN = 13.333
FRAME = (0.5, 1.5, 12.8, 7.1)


def slide_png(boxes, dpi=40):
    """A white 13.33x7.5in 'render' with a text-like texture in each box."""
    img = Image.new("L", (int(SLIDE_W_IN * dpi), int(7.5 * dpi)), 255)
    draw = ImageDraw.Draw(img)
    for left, top, width, height in boxes:
        y = top
        while y < top + height:
            draw.line([(left * dpi, y * dpi), ((left + width) * dpi, y * dpi)],
                      fill=0, width=1)
            y += 0.06
    buf = io.BytesIO()
    img.save(buf, "PNG")
    return buf.getvalue()


class TestEmptyRegion(unittest.TestCase):
    def test_a_full_frame_has_no_large_empty_region(self):
        png = slide_png([(0.5, 1.5, 12.3, 5.6)])
        region = layout_space.largest_empty_region(png, FRAME, SLIDE_W_IN)
        self.assertLess(region["share"], 0.05)

    def test_the_empty_bottom_half_is_found_and_located(self):
        png = slide_png([(0.5, 1.5, 12.3, 2.5)])
        region = layout_space.largest_empty_region(png, FRAME, SLIDE_W_IN)
        self.assertGreater(region["share"], 0.4)
        left, top, width, height = region["box_in"]
        self.assertGreater(top, 4.0)
        self.assertGreater(width, 11.0)

    def test_a_flat_card_with_little_text_counts_as_empty(self):
        img = Image.open(io.BytesIO(slide_png([(0.6, 1.6, 12.0, 0.4)])))
        draw = ImageDraw.Draw(img)
        draw.rectangle([0.5 * 40, 2.2 * 40, 12.8 * 40, 7.1 * 40], fill=240)
        buf = io.BytesIO()
        img.save(buf, "PNG")
        region = layout_space.largest_empty_region(buf.getvalue(), FRAME,
                                                   SLIDE_W_IN)
        self.assertGreater(region["share"], 0.5)


class TestEmptySpaceIssues(unittest.TestCase):
    def deck(self, count=3):
        pres = Presentation()
        pres.slide_width, pres.slide_height = Inches(13.333), Inches(7.5)
        for _ in range(count):
            pres.slides.add_slide(pres.slide_layouts[1])  # "Title and Content"
        return pres

    def test_a_half_empty_content_slide_is_reported_with_its_region(self):
        pres = self.deck()
        full = slide_png([(0.3, 1.5, 12.7, 5.8)])
        half = slide_png([(0.3, 1.5, 12.7, 2.0)])
        issues = layout_space.empty_space_issues(pres, [full, half, full],
                                                 [1, 2, 3])
        self.assertEqual([i["slide"] for i in issues], [2])
        issue = issues[0]
        self.assertEqual((issue["check"], issue["category"],
                          issue["severity"]), ("layout", "empty", "major"))
        self.assertEqual(len(issue["empty_region_in"]), 4)
        self.assertIn("left=", issue["suggested_fix"])

    def test_first_and_last_slides_are_sparse_by_design(self):
        pres = self.deck()
        blank = slide_png([])
        issues = layout_space.empty_space_issues(pres, [blank, blank, blank],
                                                 [1, 2, 3])
        self.assertEqual([i["slide"] for i in issues], [2])

    def test_can_be_switched_off(self):
        pres = self.deck()
        blank = slide_png([])
        with patch.dict(os.environ, {"VISUAL_QA_EMPTY_CHECK": "false"}):
            self.assertEqual(layout_space.empty_space_issues(
                pres, [blank] * 3, [1, 2, 3]), [])

    def test_an_unreadable_image_is_skipped_not_raised(self):
        pres = self.deck()
        self.assertEqual(layout_space.empty_space_issues(
            pres, [b"png"] * 3, [1, 2, 3]), [])

    def test_agenda_layout_is_not_treated_as_a_title(self):
        # "Agenda" contains "end"; only whole words mark a sparse layout.
        self.assertIsNone(layout_space._SPARSE_LAYOUT.search("02 Agenda"))
        self.assertIsNotNone(layout_space._SPARSE_LAYOUT.search("01 Title"))
        slide = type("S", (), {"slide_layout": type("L", (), {
            "name": "Title and Content"})()})()
        self.assertFalse(layout_space.is_sparse_by_design(slide, 2, 3))

    def test_frame_override(self):
        pres = self.deck(1)
        with patch.dict(os.environ,
                        {"VISUAL_QA_CONTENT_FRAME": "0.52,1.74,12.81,7.10"}):
            self.assertEqual(layout_space.content_frame(
                pres.slides[0], pres.slide_width, pres.slide_height),
                (0.52, 1.74, 12.81, 7.10))


def icon_png():
    buf = io.BytesIO()
    Image.new("RGBA", (40, 40), (0, 93, 185, 255)).save(buf, "PNG")
    return buf.getvalue()


class TestIconChecks(unittest.TestCase):
    def slide_with_card(self, icon_left, icon_top, card_text=None):
        pres = Presentation()
        pres.slide_width, pres.slide_height = Inches(13.333), Inches(7.5)
        slide = pres.slides.add_slide(pres.slide_layouts[6])
        card = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1),
                                      Inches(2), Inches(4), Inches(1.5))
        card.fill.solid()
        card.fill.fore_color.rgb = RGBColor(0xD5, 0x00, 0x58)
        if card_text:
            card.text_frame.text = card_text
        label = slide.shapes.add_textbox(Inches(2.2), Inches(2.3),
                                         Inches(2.6), Inches(0.4))
        label.text_frame.text = "Adjusted EBITDA €961m"
        pic = slide.shapes.add_picture(io.BytesIO(icon_png()),
                                       Inches(icon_left), Inches(icon_top),
                                       Inches(0.8), Inches(0.8))
        pic.name = layout_space.ICON_NAME_PREFIX + "euro coin"
        return pres, slide, len(slide.shapes) - 1

    def test_named_icon_concept_and_context(self):
        pres, slide, index = self.slide_with_card(1.2, 2.2)
        self.assertEqual(layout_space.icon_concept(slide.shapes[index]),
                         "euro coin")
        self.assertIn("EBITDA", layout_space.icon_context(slide, index))

    def test_a_small_square_picture_counts_as_an_icon(self):
        pres, slide, index = self.slide_with_card(1.2, 2.2)
        slide.shapes[index].name = "Picture 9"
        self.assertEqual(layout_space.icon_concept(slide.shapes[index]), "icon")

    def test_an_icon_inside_its_card_is_fine(self):
        pres, slide, _ = self.slide_with_card(1.2, 2.2)
        self.assertEqual(layout_space.icon_placement_issues(pres, [1]), [])

    def test_an_icon_crossing_the_card_edge_gets_a_spot_inside(self):
        pres, slide, index = self.slide_with_card(4.6, 2.2)
        issues = layout_space.icon_placement_issues(pres, [1])
        self.assertEqual(len(issues), 1)
        issue = issues[0]
        self.assertEqual((issue["check"], issue["severity"]),
                         ("geometry", "major"))
        self.assertIn(f"move_shape #{index} to left_in=", issue["suggested_fix"])
        left = float(issue["suggested_fix"].split("left_in=")[1].split(",")[0])
        self.assertLessEqual(left + 0.8, 5.0 - layout_space.ICON_INSET_IN + 0.01)

    def test_an_icon_drawn_over_text_inside_its_card_is_reported(self):
        # The label starts at 2.2in; an icon at (2.3, 2.3) sits on it.
        pres, slide, index = self.slide_with_card(2.3, 2.3)
        issues = layout_space.icon_placement_issues(pres, [1])
        self.assertEqual([i["category"] for i in issues], ["overlap"])
        self.assertIn("drawn over the text", issues[0]["description"])

    def test_the_outline_names_icons_and_what_they_illustrate(self):
        import deck_review
        pres, slide, index = self.slide_with_card(1.2, 2.2)
        icon = [e for e in deck_review.slide_outline(pres)[0]["elements"]
                if e["shape_index"] == index][0]
        self.assertEqual(icon["kind"], "icon")
        self.assertEqual(icon["icon_concept"], "euro coin")
        self.assertIn("EBITDA", icon["illustrates"])


class TestReplaceTextKeepingFormat(unittest.TestCase):
    def test_colour_size_and_bold_survive_a_rewrite(self):
        pres = Presentation()
        slide = pres.slides.add_slide(pres.slide_layouts[6])
        box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
        run = box.text_frame.paragraphs[0].add_run()
        run.text = "old"
        run.font.size = Pt(20)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        visual_fix.replace_text_keeping_format(box.text_frame, "new\nsecond")
        paragraphs = box.text_frame.paragraphs
        self.assertEqual([p.text for p in paragraphs], ["new", "second"])
        for para in paragraphs:
            font = para.runs[0].font
            self.assertEqual(font.size, Pt(20))
            self.assertTrue(font.bold)
            self.assertEqual(font.color.rgb, RGBColor(0xFF, 0xFF, 0xFF))


class TestZeroSizedLines(unittest.TestCase):
    def report_codes(self, pres):
        buf = io.BytesIO()
        pres.save(buf)
        return [p["code"] for p in
                deck_validation.validate_presentation(
                    pres, buf.getvalue())["problems"]]

    def test_a_horizontal_connector_is_not_zero_sized(self):
        pres = Presentation()
        slide = pres.slides.add_slide(pres.slide_layouts[6])
        slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(1),
                                   Inches(2), Inches(5), Inches(2))
        self.assertNotIn("zero_sized_shape", self.report_codes(pres))

    def test_a_zero_height_rectangle_still_is(self):
        pres = Presentation()
        slide = pres.slides.add_slide(pres.slide_layouts[6])
        slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(2),
                               Inches(4), 0)
        self.assertIn("zero_sized_shape", self.report_codes(pres))


class TestChartValues(unittest.TestCase):
    def chart(self, kind, values):
        pres = Presentation()
        slide = pres.slides.add_slide(pres.slide_layouts[6])
        return ppt_utils.add_chart(slide, kind, 1, 1, 6, 4, ["A", "B"],
                                   ["S"], [values])

    def test_column_axis_starts_at_zero(self):
        chart = self.chart("column", [4059, 4296])
        ppt_utils.style_chart_values(chart, "column", [[4059, 4296]])
        self.assertEqual(chart.value_axis.minimum_scale, 0)

    def test_line_axis_is_left_to_auto_scale(self):
        chart = self.chart("line", [21.8, 22.4])
        ppt_utils.style_chart_values(chart, "line", [[21.8, 22.4]])
        self.assertIsNone(chart.value_axis.minimum_scale)

    def test_negative_data_keeps_auto_scale(self):
        chart = self.chart("column", [-5, 10])
        ppt_utils.style_chart_values(chart, "column", [[-5, 10]])
        self.assertIsNone(chart.value_axis.minimum_scale)

    def test_data_labels_are_really_switched_on(self):
        chart = self.chart("column", [1, 2])
        ppt_utils.format_chart(chart, has_legend=False, has_data_labels=True)
        self.assertTrue(chart.plots[0].has_data_labels)
        ppt_utils.style_chart_values(chart, "column", [[1, 2]],
                                     font_size=14, number_format="#,##0",
                                     show_value_axis=False,
                                     show_gridlines=False)
        self.assertEqual(chart.plots[0].data_labels.number_format, "#,##0")
        self.assertFalse(chart.value_axis.visible)
        self.assertFalse(chart.value_axis.has_major_gridlines)
        self.assertEqual(chart.font.size, Pt(14))


class TestVisionRetry(unittest.TestCase):
    ENV = {"VISION_LLM_ENDPOINT": "https://example.invalid/responses",
           "VISION_LLM_API_KEY": "k", "VISION_LLM_MODEL": "m",
           "VISION_LLM_RETRIES": "1", "VISION_LLM_PROVIDER": "direct"}

    def test_a_timed_out_request_is_retried(self):
        ok = httpx.Response(200, json={"output_text": "{\"passed\": true}"},
                            request=httpx.Request("POST", "https://x"))
        calls = []

        def fake_post(*args, **kwargs):
            calls.append(kwargs.get("timeout"))
            if len(calls) == 1:
                raise httpx.ReadTimeout("hung")
            return ok

        with patch.dict(os.environ, self.ENV), \
                patch.object(visual_qa.httpx, "post", fake_post):
            text = visual_qa.VisionLLM().ask([], "p")
        self.assertEqual(len(calls), 2)
        self.assertIn("passed", text)

    def test_repeated_timeouts_raise_a_qa_error(self):
        def fake_post(*args, **kwargs):
            raise httpx.ReadTimeout("hung")

        with patch.dict(os.environ, self.ENV), \
                patch.object(visual_qa.httpx, "post", fake_post):
            with self.assertRaises(visual_qa.VisualQAError):
                visual_qa.VisionLLM().ask([], "p")


if __name__ == "__main__":
    unittest.main()
