"""Tests for the per-template instructions sidecar: decoding, sectioning,
attachment at template-load time, and serving."""
import base64
import io
import os
import sys
import tempfile
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from pptx import Presentation

import template_instructions
from state import PresentationStore
from tools.guidance_tools import register_guidance_tools
from tools.presentation_tools import register_presentation_tools

DOC = """# Acme deck rules

Use the template's own slides.

## 1. Slides

Slide 3 is the section divider. Never delete slide 1.

## Colours

The accent is #E4002B. Nothing else.
"""


class _FakeApp:
    """Collects the functions registered with @app.tool()."""

    def __init__(self):
        self.tools = {}

    def tool(self, **kwargs):
        def decorator(fn):
            self.tools[fn.__name__] = fn
            return fn
        return decorator


def template_bytes():
    pres = Presentation()
    buf = io.BytesIO()
    pres.save(buf)
    return buf.getvalue()


def b64(data):
    if isinstance(data, str):
        data = data.encode("utf-8")
    return base64.b64encode(data).decode("ascii")


class DecodeTests(unittest.TestCase):
    def tearDown(self):
        os.environ.pop("TEMPLATE_INSTRUCTIONS_MAX_KB", None)

    def test_plain_markdown_passes_through(self):
        self.assertEqual(template_instructions.decode_payload(DOC), DOC)

    def test_base64_payload(self):
        self.assertEqual(template_instructions.decode_payload(b64(DOC)), DOC)

    def test_data_uri(self):
        uri = "data:text/markdown;base64," + b64(DOC)
        self.assertEqual(template_instructions.decode_payload(uri), DOC)

    def test_absent_payloads_are_not_errors(self):
        for payload in (None, "", "   ",
                        "file:data::files/bucket/deck.md",
                        "files/bucket/missing.md"):
            self.assertIsNone(template_instructions.decode_payload(payload),
                              payload)

    def test_binary_is_rejected(self):
        with self.assertRaises(template_instructions.InstructionsError) as ctx:
            template_instructions.decode_payload(b64(template_bytes()))
        self.assertIn("not a markdown document", str(ctx.exception))

    def test_size_cap(self):
        os.environ["TEMPLATE_INSTRUCTIONS_MAX_KB"] = "1"
        with self.assertRaises(template_instructions.InstructionsError) as ctx:
            template_instructions.decode_payload("x" * 4096)
        self.assertIn("limit", str(ctx.exception))

    def test_invalid_cap_falls_back(self):
        os.environ["TEMPLATE_INSTRUCTIONS_MAX_KB"] = "lots"
        self.assertEqual(template_instructions.max_bytes(),
                         int(template_instructions.DEFAULT_MAX_KB * 1024))


class ParseTests(unittest.TestCase):
    def test_numbered_and_unnumbered_sections(self):
        doc = template_instructions.parse(DOC)
        self.assertEqual([s["section"]
                          for s in template_instructions.section_list(doc)],
                         ["slides", "colours"])

    def test_section_body_is_scoped(self):
        doc = template_instructions.parse(DOC)
        picked = template_instructions.select(doc, "colours")
        self.assertIn("#E4002B", picked["instructions"])
        self.assertNotIn("section divider", picked["instructions"])

    def test_substring_match(self):
        doc = template_instructions.parse(DOC)
        self.assertEqual(template_instructions.select(doc, "colour")["section"],
                         "colours")

    def test_unknown_section_lists_available(self):
        doc = template_instructions.parse(DOC)
        picked = template_instructions.select(doc, "fonts")
        self.assertIn("error", picked)
        self.assertIn("colours", picked["error"])

    def test_document_without_headings_serves_whole(self):
        doc = template_instructions.parse("Just one rule: keep it blue.")
        self.assertEqual(doc["sections"], {})
        self.assertIn("keep it blue",
                      template_instructions.select(doc, None)["instructions"])
        self.assertIn("error", template_instructions.select(doc, "colours"))


class LoadAndServeTests(unittest.TestCase):
    def setUp(self):
        self.store = PresentationStore(ttl_seconds=60, max_items=10)
        app = _FakeApp()
        register_presentation_tools(app, self.store,
                                    lambda: None, lambda: [])
        register_guidance_tools(app, self.store)
        self.create = app.tools["create_presentation_from_template_content"]
        self.create_from_path = app.tools["create_presentation_from_template"]
        self.get = app.tools["get_template_instructions"]

    def test_instructions_attached_in_the_same_call(self):
        result = self.create(b64(template_bytes()), b64(DOC))
        info = result["template_instructions"]
        self.assertTrue(info["loaded"])
        self.assertEqual([s["section"] for s in info["sections"]],
                         ["slides", "colours"])
        # The summary reports the document without spending it.
        self.assertNotIn("E4002B", str(info))

        served = self.get(result["presentation_id"], "slides")
        self.assertIn("section divider", served["instructions"])
        self.assertIn("trust", served)

    def test_missing_sidecar_still_loads_the_template(self):
        for payload in (None, "file:data::files/bucket/missing.md"):
            result = self.create(b64(template_bytes()), payload)
            self.assertIn("presentation_id", result)
            info = result["template_instructions"]
            self.assertFalse(info["loaded"])
            self.assertEqual(info["reason"], "no_instructions_supplied")
            self.assertIn("error", self.get(result["presentation_id"]))

    def test_broken_sidecar_is_reported_not_fatal(self):
        result = self.create(b64(template_bytes()), b64(template_bytes()))
        self.assertIn("presentation_id", result)
        info = result["template_instructions"]
        self.assertFalse(info["loaded"])
        self.assertIn("not a markdown document", info["reason"])

    def test_unknown_presentation_id(self):
        self.assertIn("Unknown or expired", self.get("nope")["error"])

    def test_sidecar_found_beside_a_template_on_disk(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, "deck.pptx")
            Presentation().save(path)
            Path(tmp, "deck.md").write_text(DOC, encoding="utf-8")
            result = self.create_from_path(path)
            info = result["template_instructions"]
            self.assertTrue(info["loaded"])
            self.assertEqual(info["path"], os.path.join(tmp, "deck.md"))
            self.assertIn("#E4002B",
                          self.get(result["presentation_id"],
                                   "colours")["instructions"])

    def test_no_sidecar_beside_a_template_on_disk(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, "deck.pptx")
            Presentation().save(path)
            result = self.create_from_path(path)
            self.assertFalse(result["template_instructions"]["loaded"])

    def test_instructions_die_with_the_deck(self):
        result = self.create(b64(template_bytes()), b64(DOC))
        pid = result["presentation_id"]
        del self.store[pid]
        self.assertIsNone(self.store.get_instructions(pid))


if __name__ == "__main__":
    unittest.main()
