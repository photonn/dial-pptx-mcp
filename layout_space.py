"""Deterministic empty-space check for rendered slides.

The vision reviewers are good at "this text overflows" and bad at "half of this
slide is unused": across whole-deck reviews of decks whose content slides were
40% empty they reported overflow and legibility and never once the empty
area, which is the defect a reader notices first. Emptiness is also something
pixels answer exactly, so it is measured here rather than asked for.

A slide's *content frame* is the area between its title (and subline) and its
footer band, inside the title's side margins — derived from the layout's
placeholders, or given explicitly with VISUAL_QA_CONTENT_FRAME. The rendered
frame is cut into ~0.1in cells; a cell holds *ink* when its pixels vary (text,
lines, chart marks, picture detail, edges of shapes). Ink is dilated by
INK_MARGIN_IN so ordinary breathing room around content does not count, and
the largest ink-free rectangle is what gets reported: a flat card interior
under a two-line text is as empty as bare slide background.

Title, divider, statement and closing slides are sparse by design and skipped.
"""
import io
import os
import re

from PIL import Image, ImageStat
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.ns import qn

from logging_utils import get_logger

logger = get_logger("layout_space")

# add_icon_to_slide names each placed icon "Icon: <concept>"; that name is how
# an icon is told apart from a photo, and what it is meant to show.
ICON_NAME_PREFIX = "Icon: "
# An icon inside a card keeps at least this much of the card around it.
ICON_INSET_IN = 0.12
# Average glyph width of mixed-case sans-serif text, in ems — enough to tell
# where a line of text ends, which is all the icon check needs.
CHAR_WIDTH_EM = 0.48

CELL_IN = 0.1
# Uniform-looking cells vary by JPEG-free renderer noise only; text strokes and
# shape edges push the standard deviation far above this.
INK_STDDEV = 4.0
# Space around content that is just breathing room, not emptiness.
INK_MARGIN_IN = 0.3
ANALYSIS_DPI = 40
# Largest empty rectangle, as a share of the content frame, that is reported.
# Calibrated on hand-built reference decks, whose content slides stay under
# 0.10, against generated ones whose half-empty slides score 0.3-0.6.
DEFAULT_THRESHOLD = 0.10
MIN_SIDE_IN = 1.0

_SPARSE_LAYOUT = re.compile(
    r"\b(title|closing|end|divider|section|chapter|statement|quote|thank)",
    re.IGNORECASE)


def threshold():
    try:
        return float(os.environ.get("VISUAL_QA_EMPTY_THRESHOLD",
                                    DEFAULT_THRESHOLD))
    except ValueError:
        return DEFAULT_THRESHOLD


def enabled():
    return os.environ.get("VISUAL_QA_EMPTY_CHECK", "true").strip().lower() \
        not in ("0", "false", "no", "off")


def _emu_in(value):
    return value / 914400


def _frame_override():
    raw = os.environ.get("VISUAL_QA_CONTENT_FRAME", "").strip()
    if not raw:
        return None
    try:
        left, top, right, bottom = (float(v) for v in raw.split(","))
    except ValueError:
        logger.warning("config_invalid var=VISUAL_QA_CONTENT_FRAME value=%s",
                       raw)
        return None
    return left, top, right, bottom


def content_frame(slide, slide_width, slide_height):
    """(left, top, right, bottom) in inches of the area content should use,
    or None when the slide's layout gives no title to measure from."""
    override = _frame_override()
    if override:
        return override
    width_in, height_in = _emu_in(slide_width), _emu_in(slide_height)
    title = None
    top_band, footer_top = [], height_in
    for ph in slide.slide_layout.placeholders:
        if ph.top is None or ph.height is None:
            continue
        top, bottom = _emu_in(ph.top), _emu_in(ph.top + ph.height)
        kind = str(ph.placeholder_format.type or "")
        if "TITLE" in kind and title is None:
            title = ph
            top_band.append(bottom)
        elif bottom < height_in * 0.3 and top < height_in * 0.25:
            # a subline under the title
            top_band.append(bottom)
        elif top > height_in * 0.85:
            footer_top = min(footer_top, top)
    if title is None or title.left is None:
        return None
    left = _emu_in(title.left)
    top = max(top_band) + 0.15
    bottom = footer_top - 0.05
    right = width_in - left
    if bottom - top < 1.0 or right - left < 2.0:
        return None
    return left, top, right, bottom


# ...unless the name says the layout carries content ("Title and Content").
_CONTENT_LAYOUT = re.compile(
    r"content|text|chart|table|two|comparison|canvas|picture|caption|agenda",
    re.IGNORECASE)


def is_sparse_by_design(slide, number, total):
    if number in (1, total):
        return True
    name = slide.slide_layout.name or ""
    return bool(_SPARSE_LAYOUT.search(name)) \
        and not _CONTENT_LAYOUT.search(name)


def _largest_rect(grid):
    """Largest all-True rectangle in a list-of-lists bool grid ->
    (cells, (row0, col0, row1, col1))."""
    rows = len(grid)
    cols = len(grid[0]) if rows else 0
    heights = [0] * cols
    best = (0, None)
    for r in range(rows):
        heights = [heights[c] + 1 if grid[r][c] else 0 for c in range(cols)]
        stack = []
        for c in range(cols + 1):
            current = heights[c] if c < cols else 0
            start = c
            while stack and stack[-1][1] >= current:
                start, height = stack.pop()
                if height * (c - start) > best[0]:
                    best = (height * (c - start),
                            (r - height + 1, start, r, c - 1))
            stack.append((start, current))
    return best


def largest_empty_region(png, frame, slide_width_in):
    """-> {"share", "box_in": [left, top, width, height]} for the largest
    ink-free rectangle inside frame, or None when the frame is unusable."""
    image = Image.open(io.BytesIO(png)).convert("L")
    scale = ANALYSIS_DPI / (image.width / slide_width_in)
    if scale < 1:
        image = image.resize((max(1, int(image.width * scale)),
                              max(1, int(image.height * scale))))
    ppi = image.width / slide_width_in
    left, top, right, bottom = frame
    x0, y0 = int(left * ppi), int(top * ppi)
    x1, y1 = min(image.width, int(right * ppi)), min(image.height,
                                                    int(bottom * ppi))
    cell = max(2, int(round(CELL_IN * ppi)))
    cols, rows = (x1 - x0) // cell, (y1 - y0) // cell
    if cols < 3 or rows < 3:
        return None
    ink = [[ImageStat.Stat(image.crop((x0 + c * cell, y0 + r * cell,
                                       x0 + (c + 1) * cell,
                                       y0 + (r + 1) * cell))).stddev[0]
            >= INK_STDDEV for c in range(cols)] for r in range(rows)]
    reach = max(1, int(round(INK_MARGIN_IN / (cell / ppi))))
    used = [[False] * cols for _ in range(rows)]
    for r in range(rows):
        for c in range(cols):
            if ink[r][c]:
                for rr in range(max(0, r - reach), min(rows, r + reach + 1)):
                    row = used[rr]
                    for cc in range(max(0, c - reach),
                                    min(cols, c + reach + 1)):
                        row[cc] = True
    area, box = _largest_rect([[not u for u in row] for row in used])
    if not box:
        return {"share": 0.0, "box_in": None}
    cell_in = cell / ppi
    r0, c0, r1, c1 = box
    return {"share": round(area / (rows * cols), 3),
            "box_in": [round(left + c0 * cell_in, 2),
                       round(top + r0 * cell_in, 2),
                       round((c1 - c0 + 1) * cell_in, 2),
                       round((r1 - r0 + 1) * cell_in, 2)]}


def empty_space_issues(pres, images, image_slides):
    """Blocking "empty" issues for content slides whose largest unused area
    exceeds the threshold. Each carries the region in inches so the author
    can build into exactly that space."""
    if not enabled():
        return []
    issues = []
    total = len(pres.slides)
    width_in = _emu_in(pres.slide_width)
    limit = threshold()
    for png, number in zip(images, image_slides):
        slide = pres.slides[number - 1]
        if is_sparse_by_design(slide, number, total):
            continue
        frame = content_frame(slide, pres.slide_width, pres.slide_height)
        if frame is None:
            continue
        try:
            region = largest_empty_region(png, frame, width_in)
        except Exception as e:  # a measurement must never fail the review
            logger.warning("empty_check_failed slide=%d error=%s", number, e)
            continue
        if not region or not region["box_in"]:
            continue
        x, y, w, h = region["box_in"]
        logger.debug("empty_check slide=%d share=%.3f box=%s", number,
                     region["share"], region["box_in"])
        if region["share"] < limit or w < MIN_SIDE_IN or h < MIN_SIDE_IN:
            continue
        issues.append({
            "slide": number,
            "severity": "major",
            "category": "empty",
            "check": "layout",
            "element": "content area",
            "empty_region_in": region["box_in"],
            "description": (
                f"{round(region['share'] * 100)}% of the content area is "
                f"unused: nothing is drawn in the {w}x{h} in region at "
                f"left={x}, top={y} (a card or panel with nothing inside "
                f"counts as unused)."),
            "suggested_fix": (
                f"Fill the region left={x}, top={y}, width={w}, height={h} "
                f"(inches) with content that supports the headline — a "
                f"chart, KPI tiles, icon cards, a callout panel — or enlarge "
                f"the neighbouring elements into it; if a card is mostly "
                f"empty, add its icon/figure or shrink it and use the space."),
        })
    return issues


# ---------------------------------------------------------------- icons

def icon_concept(shape):
    """The concept of an icon placed by add_icon_to_slide, else None."""
    if shape.shape_type != MSO_SHAPE_TYPE.PICTURE:
        return None
    name = getattr(shape, "name", "") or ""
    if name.startswith(ICON_NAME_PREFIX):
        return name[len(ICON_NAME_PREFIX):].strip() or "icon"
    # A small square picture is an icon too (a catalog PNG, or one placed
    # before icons were named) — its concept is just unknown.
    box = _box_in(shape)
    if box and box[2] <= 1.3 and box[3] <= 1.3 and \
            abs(box[2] - box[3]) <= 0.05 * max(box[2], box[3]):
        return "icon"
    return None


def _box_in(shape):
    if None in (shape.left, shape.top, shape.width, shape.height):
        return None
    return (_emu_in(shape.left), _emu_in(shape.top),
            _emu_in(shape.width), _emu_in(shape.height))


def _overlap(a, b):
    width = min(a[0] + a[2], b[0] + b[2]) - max(a[0], b[0])
    height = min(a[1] + a[3], b[1] + b[3]) - max(a[1], b[1])
    return max(0.0, width) * max(0.0, height)


def _is_container(shape):
    """A filled autoshape big enough to hold an icon: card, tile, panel, bar."""
    if shape.shape_type != MSO_SHAPE_TYPE.AUTO_SHAPE:
        return False
    try:
        if shape.fill.type != 1:  # MSO_FILL.SOLID
            return False
    except Exception:
        return False
    box = _box_in(shape)
    return bool(box) and box[2] >= 0.5 and box[3] >= 0.5


def icon_context(slide, icon_index):
    """Text the icon at icon_index illustrates: the text of the smallest
    container holding its centre (a card with its own text), else of the
    nearest text-bearing shape to its right or below — the label an icon
    normally sits beside."""
    shapes = list(slide.shapes)
    icon_box = _box_in(shapes[icon_index])
    if not icon_box:
        return ""
    cx, cy = icon_box[0] + icon_box[2] / 2, icon_box[1] + icon_box[3] / 2
    texts = []
    holder = None
    for shape in shapes:
        box = _box_in(shape)
        if not box or shape is shapes[icon_index]:
            continue
        inside = box[0] <= cx <= box[0] + box[2] and \
            box[1] <= cy <= box[1] + box[3]
        if inside and _is_container(shape) and (
                holder is None or box[2] * box[3] < holder[1]):
            holder = (shape, box[2] * box[3], box)
    region = holder[2] if holder else (icon_box[0] - 0.2, icon_box[1] - 0.3,
                                       icon_box[2] + 4.0, icon_box[3] + 0.8)
    for shape in shapes:
        if not shape.has_text_frame or not shape.text_frame.text.strip():
            continue
        if getattr(shape, "is_placeholder", False):
            try:
                if "TITLE" in str(shape.placeholder_format.type):
                    continue
            except Exception:
                pass
        box = _box_in(shape)
        if box and _overlap(box, region) > 0.3 * box[2] * box[3]:
            texts.append(" ".join(shape.text_frame.text.split()))
    return " / ".join(texts)[:160]


def _text_extent(shape):
    """Estimated box of the text a shape actually shows (inches), not of the
    whole shape: a card's text box often spans the card while its two lines
    use the top third, and an icon below them is not "on the text"."""
    if not getattr(shape, "has_text_frame", False):
        return None
    frame = shape.text_frame
    paragraphs = [p for p in frame.paragraphs if p.text.strip()]
    if not paragraphs:
        return None
    box = _box_in(shape)
    if not box:
        return None
    base = _layout_placeholder(shape)
    size_pt = _inherited_size_pt(base) or 14.0
    for para in paragraphs:
        for run in para.runs:
            if run.font.size:
                size_pt = run.font.size.pt
                break
        else:
            continue
        break
    inner_w = max(0.3, box[2] - 0.2)
    chars_per_line = max(1, int(inner_w / (size_pt * CHAR_WIDTH_EM / 72)))
    lines = sum(max(1, -(-len(p.text) // chars_per_line)) for p in paragraphs)
    height = min(box[3], lines * size_pt * 1.25 / 72 + 0.1)
    anchor = frame.vertical_anchor
    if anchor is None and base is not None:
        # A placeholder inherits its anchor from the layout: the "so what"
        # panel of a split layout centres its text, and reading it as
        # top-anchored put every icon above that text "on" it.
        try:
            anchor = base.text_frame.vertical_anchor
        except Exception:
            anchor = None
    anchor = str(anchor or "")
    if "MIDDLE" in anchor:
        top = box[1] + (box[3] - height) / 2
    elif "BOTTOM" in anchor:
        top = box[1] + box[3] - height
    else:
        top = box[1]
    left, width = box[0], box[2]
    if lines == len(paragraphs):
        # Every paragraph fits one line: the text is only as wide as its
        # longest line, placed by its alignment.
        width = min(box[2], max(len(p.text) for p in paragraphs)
                    * size_pt * CHAR_WIDTH_EM / 72 + 0.2)
        align = str(paragraphs[0].alignment or "")
        if "CENTER" in align:
            left = box[0] + (box[2] - width) / 2
        elif "RIGHT" in align:
            left = box[0] + box[2] - width
    return (left, top, width, height)


def _layout_placeholder(shape):
    if not getattr(shape, "is_placeholder", False):
        return None
    try:
        return shape._base_placeholder
    except Exception:
        return None


def _inherited_size_pt(base):
    """First explicit font size in the layout placeholder's own list style
    or text, if any."""
    if base is None:
        return None
    try:
        for node in base._element.iter(qn("a:defRPr"), qn("a:rPr")):
            if node.get("sz"):
                return int(node.get("sz")) / 100
    except Exception:
        return None
    return None


def _free_spot(icon_box, container, text_boxes, shrink=False):
    """A position (left, top, size) inside container, clear of text, for an
    icon that crosses its edge: where it is (pulled inside), then the
    container's corners and side middles; with shrink, smaller sizes too."""
    cl, ct, cw, ch = container
    inset = ICON_INSET_IN
    size = icon_box[2]
    while size >= 0.45:
        lo_x, hi_x = cl + inset, cl + cw - size - inset
        lo_y, hi_y = ct + inset, ct + ch - size - inset
        if hi_x >= lo_x and hi_y >= lo_y:
            mid_y = ct + (ch - size) / 2
            candidates = [
                (min(max(icon_box[0], lo_x), hi_x),
                 min(max(icon_box[1], lo_y), hi_y)),
                (lo_x, lo_y), (hi_x, lo_y), (lo_x, mid_y), (hi_x, mid_y),
                (hi_x, hi_y), (lo_x, hi_y)]
            for left, top in candidates:
                spot = (left, top, size, size)
                if not any(_overlap(spot, t) > 0.02 for t in text_boxes
                           if _overlap(t, container) > 0):
                    return round(left, 2), round(top, 2), round(size, 2)
        if not shrink:
            break
        size -= 0.15
    return None


def _make_room(icon_box, container, shapes, container_index):
    """When no text-free spot exists: put the icon in the container's
    top-left corner and move the one text box starting at the container's
    top down below it — if the text still fits inside the container.
    -> (text shape index, new text top, icon left, icon top) or None."""
    cl, ct, cw, ch = container
    size = icon_box[2]
    left, top = cl + ICON_INSET_IN, ct + ICON_INSET_IN
    for index, shape in enumerate(shapes):
        if index == container_index:
            continue  # the card's own text cannot move without the card
        extent = _text_extent(shape)
        box = _box_in(shape)
        if not extent or not box or _overlap(box, container) < 0.8 * box[2] * box[3]:
            continue
        if extent[1] > top + size:
            continue  # already below where the icon goes
        new_top = top + size + 0.08
        if new_top + extent[3] <= ct + ch - 0.05:
            return index, round(new_top, 2), round(left, 2), round(top, 2)
    return None


def covered_text(icon_index, icon_box, shapes):
    """Index of a text shape whose estimated text the icon covers by more
    than a quarter of the icon's area, else None."""
    area = icon_box[2] * icon_box[3]
    for index, shape in enumerate(shapes):
        if index == icon_index:
            continue
        extent = _text_extent(shape)
        if extent and _overlap(icon_box, extent) > 0.25 * area:
            return index
    return None


def _holder(icon_box, containers):
    """The smallest container holding the icon's centre, as (index, box)."""
    cx = icon_box[0] + icon_box[2] / 2
    cy = icon_box[1] + icon_box[3] / 2
    best = None
    for index, box in containers:
        if box[0] <= cx <= box[0] + box[2] and box[1] <= cy <= box[1] + box[3]:
            if best is None or box[2] * box[3] < best[1][2] * best[1][3]:
                best = (index, box)
    return best


def icon_placement_issues(pres, slide_numbers):
    """Geometric faults of placed icons, each with the position that fixes
    it — so the repair planner can simply move the icon there:

    - straddling: the icon crosses the edge of a card, tile, bar or panel it
      mostly sits on (half inside a magenta bar, over a blue panel's top);
    - off frame: the icon leaves the slide.
    """
    issues = []
    width_in = _emu_in(pres.slide_width)
    height_in = _emu_in(pres.slide_height)
    for number in slide_numbers:
        slide = pres.slides[number - 1]
        shapes = list(slide.shapes)
        containers = [(i, _box_in(s)) for i, s in enumerate(shapes)
                      if _is_container(s)]
        text_boxes = [b for b in (_text_extent(sh) for sh in shapes) if b]
        frame = None
        if not is_sparse_by_design(slide, number, len(pres.slides)):
            frame = content_frame(slide, pres.slide_width, pres.slide_height)
        if frame:
            frame = (frame[0], frame[1], min(frame[2], width_in),
                     min(frame[3], height_in))
        for index, shape in enumerate(shapes):
            concept = icon_concept(shape)
            if concept is None:
                continue
            box = _box_in(shape)
            if not box:
                continue
            area = box[2] * box[3]
            target = None
            for c_index, c_box in containers:
                share = _overlap(box, c_box) / area if area else 0
                if not 0.05 < share < 0.97:
                    continue
                if c_box[2] < box[2] + 2 * ICON_INSET_IN or \
                        c_box[3] < box[3] + 2 * ICON_INSET_IN:
                    continue
                target = (c_index, share,
                          _free_spot(box, c_box, text_boxes))
                break
            if target:
                c_index, share, spot = target
                c_box = containers[[c for c, _ in containers].index(c_index)][1]
                # Full size where it fits; else push the text down; shrinking
                # the icon (and breaking the row's consistency) comes last.
                shift = None if spot else _make_room(box, c_box, shapes,
                                                     c_index)
                if not spot and not shift:
                    spot = _free_spot(box, c_box, text_boxes, shrink=True)
                if shift:
                    text_index, text_top, left, top = shift
                    fix = (f"move_shape #{text_index} (the card's text) down "
                           f"to top_in={text_top}, then move_shape #{index} "
                           f"to left_in={left}, top_in={top}: the icon then "
                           f"sits inside #{c_index} above its text.")
                elif spot:
                    left, top, size = spot
                    fix = (f"move_shape #{index} to left_in={left}, "
                           f"top_in={top}" + (
                               f" and resize_shape it to width_in={size}, "
                               f"height_in={size}" if size < box[2] - 0.01
                               else "") +
                           f": that spot is inside #{c_index} and clear of "
                           f"its text.")
                else:
                    fix = (f"there is no text-free spot for the icon inside "
                           f"#{c_index}: delete icon #{index}, or move the "
                           f"card's text to make room and move the icon "
                           f"fully inside #{c_index}.")
                issues.append({
                    "slide": number, "severity": "major",
                    "category": "layout", "check": "geometry",
                    "element": f"#{index} icon ({concept})",
                    "description": (
                        f"Icon #{index} ({concept}) crosses the edge of "
                        f"shape #{c_index}: {round((1 - share) * 100)}% of "
                        f"it hangs outside the card/panel it sits on."),
                    "suggested_fix": fix,
                })
            elif covered_text(index, box, shapes):
                # Fully inside its card but drawn over text — typically a
                # card that carries its own text with icons placed on it.
                t_index = covered_text(index, box, shapes)
                holder = _holder(box, containers)
                spot = None
                if holder:
                    spot = _free_spot(box, holder[1], text_boxes) or \
                        _free_spot(box, holder[1], text_boxes, shrink=True)
                if spot:
                    left, top, size = spot
                    fix = (f"move_shape #{index} to left_in={left}, "
                           f"top_in={top}" + (
                               f" and resize_shape it to width_in={size}, "
                               f"height_in={size}" if size < box[2] - 0.01
                               else "") + ": clear of the text.")
                else:
                    fix = (f"delete icon #{index}, or move the text of "
                           f"#{t_index} clear of it.")
                issues.append({
                    "slide": number, "severity": "major",
                    "category": "overlap", "check": "geometry",
                    "element": f"#{index} icon ({concept})",
                    "description": (f"Icon #{index} ({concept}) is drawn over "
                                    f"the text of shape #{t_index}."),
                    "suggested_fix": fix,
                })
            elif frame and (box[0] < frame[0] - 0.05
                            or box[0] + box[2] > frame[2] + 0.05
                            or box[1] + box[3] > frame[3] + 0.05):
                # Past the content frame's sides or bottom (the top band is
                # the title's, and a header icon may legitimately sit there).
                left = round(min(max(box[0], frame[0]),
                                 frame[2] - box[2]), 2)
                top = round(min(box[1], frame[3] - box[3]), 2)
                issues.append({
                    "slide": number, "severity": "major",
                    "category": "layout", "check": "geometry",
                    "element": f"#{index} icon ({concept})",
                    "description": (f"Icon #{index} ({concept}) sticks out of "
                                    f"the content area."),
                    "suggested_fix": f"move_shape #{index} to left_in={left}, "
                                     f"top_in={top}.",
                })
    return issues
